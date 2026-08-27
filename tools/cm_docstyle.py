"""Word-like Excel layout: no gridlines, headings, callouts, hyperlinks."""

import struct

from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.page import PageMargins

NAVY = "1F4E79"
TEAL = "2E75B6"
TEXT = "2F2F2F"
MUTED = "5B6770"
LINE = "C5CDD4"
HAIR = "E6E9ED"
CORE_BG = "FFF6E8"
COND_BG = "EEF5FA"
CALC_BG = "F4F1EA"
HDR_BG = "EEF1F4"
ADV_BG = "F3F8F3"
LIM_BG = "FBF4F2"
LINK = "0563C1"
WHITE = "FFFFFF"
GOLD = "8A6D3B"
VALUE_BLUE = "0070C0"

COLS = 9
W = [5, 16, 26, 14, 26, 8, 14, 34, 38]

thin = Border(
    left=Side(style="hair", color=HAIR),
    right=Side(style="hair", color=HAIR),
    top=Side(style="hair", color=HAIR),
    bottom=Side(style="hair", color=HAIR),
)
BOX_LINE = "000000"
box_side = Side(style="thin", color=BOX_LINE)
bottom_rule = Border(bottom=Side(style="medium", color=NAVY))
no_b = Border()


def fl(h):
    return PatternFill("solid", fgColor=h)


def ft(size=11, bold=False, color=TEXT, italic=False, underline=None):
    return Font(name="Calibri", size=size, bold=bold, color=color, italic=italic, underline=underline)


L = Alignment(wrap_text=True, vertical="center", horizontal="left")
T = Alignment(wrap_text=True, vertical="top", horizontal="left")
C = Alignment(wrap_text=True, vertical="center", horizontal="center")


class DocSheet:
    def __init__(self, ws, footer=""):
        self.ws = ws
        self.r = 1
        for i, w in enumerate(W, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins = PageMargins(0.55, 0.55, 0.6, 0.55)
        ws.oddHeader.left.text = "Mobility Management in Connected Mode  |  Huawei eRAN21.1"
        ws.oddFooter.left.text = footer
        ws.oddFooter.right.text = "Page &P of &N"
        ws.sheet_format.defaultRowHeight = 16
        ws.freeze_panes = "A6"
        ws.print_title_rows = "1:5"
        ws.sheet_properties.tabColor = NAVY

    def _merge(self, c1, c2, value, size=11, bold=False, color=TEXT, bg=None, align=None, h=None, italic=False, border=None):
        ws, r = self.ws, self.r
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        cell = ws.cell(r, c1, value)
        cell.font = ft(size, bold, color, italic)
        cell.alignment = align or T
        if bg:
            fill = fl(bg)
            cell.fill = fill
            for c in range(c1 + 1, c2 + 1):
                ws.cell(r, c).fill = fill
                ws.cell(r, c).border = border or no_b
        cell.border = border or no_b
        if h:
            ws.row_dimensions[r].height = h
        return cell

    def space(self, h=10):
        self.ws.row_dimensions[self.r].height = h
        self.r += 1
        return self

    def banner(self, text):
        self._merge(1, COLS, "  " + text, size=10, bold=True, color=WHITE, bg=NAVY, align=L, h=22)
        self.r += 1
        return self

    def title(self, chapter, name):
        self._merge(1, COLS, f"  {chapter}   {name}", size=20, bold=True, color=NAVY, align=L, h=34)
        self.r += 1
        return self

    def meta(self, text):
        self._merge(1, COLS, "  " + text, size=10, color=MUTED, italic=True, align=L, h=20)
        self.r += 1
        return self

    def nav(self, items):
        """items: list of (label, sheet_name or None)."""
        ws, r = self.ws, self.r
        ws.row_dimensions[r].height = 20
        used = min(len(items), COLS)
        for i, (label, sheet) in enumerate(items[:COLS]):
            cell = ws.cell(r, i + 1, label)
            cell.alignment = L
            cell.border = no_b
            if sheet:
                cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1", display=label)
                cell.font = ft(10, True, LINK, underline="single")
            else:
                cell.font = ft(10, False, MUTED)
        self.r += 1
        return self

    def h1(self, text):
        self.space(14)
        self._merge(1, COLS, text, size=13, bold=True, color=NAVY, align=L, h=24, border=bottom_rule)
        self.r += 1
        self.space(6)
        return self

    def h2(self, text):
        self.space(10)
        self._merge(1, COLS, text, size=12, bold=True, color=TEAL, align=L, h=22)
        self.r += 1
        self.space(4)
        return self

    def para(self, text, h=None):
        n = max(1, (len(text) + 95) // 96) + text.count("\n")
        self._merge(1, COLS, text, size=11, color=TEXT, align=T, h=h or min(90, 16 + n * 14))
        self.r += 1
        return self

    def bullets(self, items):
        for it in items:
            self._merge(1, COLS, "    •  " + it, size=11, color=TEXT, align=T, h=min(120, 22 + (len(it) // 75) * 14))
            self.r += 1
        return self

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            self._merge(1, COLS, f"    {i}.  " + it, size=11, color=TEXT, align=T, h=min(120, 22 + (len(it) // 75) * 14))
            self.r += 1
        return self

    def callout(self, kind, title, lines):
        bg = {"CORE": CORE_BG, "CONDITION": COND_BG, "CALC": CALC_BG}.get(kind, HDR_BG)
        tc = {"CORE": GOLD, "CONDITION": TEAL, "CALC": NAVY}.get(kind, NAVY)
        self.space(6)
        start = self.r
        self._merge(1, COLS, "  " + title, size=10, bold=True, color=tc, bg=bg, align=L, h=18)
        self.r += 1
        body = lines if isinstance(lines, str) else "\n".join(lines)
        wrap_extra = sum(max(0, (len(line) - 155) // 155) for line in body.split("\n"))
        n = body.count("\n") + 1 + wrap_extra
        self._merge(1, COLS, body, size=11, color=TEXT, bg=bg, align=T, h=min(420, max(36, 16 + n * 16)))
        self.r += 1
        self._paint_range(start, 1, self.r - 1, COLS)
        for c in range(1, COLS + 1):
            self.ws.cell(start, c).fill = fl(bg)
            if self.r - 1 > start:
                self.ws.cell(self.r - 1, c).fill = fl(bg)
        return self

    def two_col(self, left_title, left_items, right_title, right_items):
        ws, r = self.ws, self.r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=COLS)
        a = ws.cell(r, 1, "  Advantage")
        b = ws.cell(r, 4, "  Limitation")
        a.font = b.font = ft(10, True, NAVY)
        a.fill = fl(ADV_BG)
        b.fill = fl(LIM_BG)
        a.alignment = b.alignment = L
        ws.row_dimensions[r].height = 18
        for c in range(1, 4):
            ws.cell(r, c).fill = fl(ADV_BG)
            ws.cell(r, c).border = no_b
        for c in range(4, COLS + 1):
            ws.cell(r, c).fill = fl(HDR_BG if c == 4 else LIM_BG)
            ws.cell(r, c).fill = fl(LIM_BG)
            ws.cell(r, c).border = no_b
        a.fill = fl(ADV_BG)
        hdr = r
        self.r += 1
        n = max(len(left_items), len(right_items), 1)
        left = left_items + [""] * (n - len(left_items))
        right = right_items + [""] * (n - len(right_items))
        body_start = self.r
        for i in range(n):
            r = self.r
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=COLS)
            lv = ("    •  " + left[i]) if left[i] else ""
            rv = ("    •  " + right[i]) if right[i] else ""
            la = ws.cell(r, 1, lv)
            ra = ws.cell(r, 4, rv)
            la.font = ra.font = ft(11, False, TEXT)
            la.alignment = ra.alignment = T
            h = min(64, 20 + max(len(left[i]), len(right[i])) // 50 * 12)
            ws.row_dimensions[r].height = h
            for c in range(1, 4):
                ws.cell(r, c).fill = fl(ADV_BG)
            for c in range(4, COLS + 1):
                ws.cell(r, c).fill = fl(LIM_BG)
            self.r += 1
        last = self.r - 1
        self._paint_range(hdr, 1, last, COLS)
        for c in range(1, 4):
            ws.cell(hdr, c).fill = fl(ADV_BG)
        for c in range(4, COLS + 1):
            ws.cell(hdr, c).fill = fl(LIM_BG)
        for r in range(body_start, last + 1):
            for c in range(1, 4):
                ws.cell(r, c).fill = fl(ADV_BG)
            for c in range(4, COLS + 1):
                ws.cell(r, c).fill = fl(LIM_BG)
        return self

    def param_heads(self):
        names = ["SN", "MO", "Parameter", "Value", "Comment", "Core", "Sub-group", "Parameter Meaning", "Reference"]
        ws, r = self.ws, self.r
        ws.row_dimensions[r].height = 22
        for c, name in enumerate(names, 1):
            cell = ws.cell(r, c, name)
            cell.font = ft(9, True, NAVY)
            cell.fill = fl(HDR_BG)
            cell.alignment = C
            cell.border = Border(bottom=Side(style="thin", color=NAVY))
        self.r += 1
        return self

    def param_row(self, sn, mo, param, value, comment, core, group, meaning="", reference="", link_sheet=None):
        ws, r = self.ws, self.r
        vals = [sn, mo, param, value, comment, core, group, meaning, reference]
        h = min(84, 22 + max(len(str(comment)), len(str(meaning)), len(str(reference)), 36) // 36 * 12)
        ws.row_dimensions[r].height = h
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            if c == 4:
                cell.font = ft(10, True, VALUE_BLUE)
            elif c in (1, 6):
                cell.font = ft(10, True, TEXT)
            else:
                cell.font = ft(10, False, TEXT)
            cell.alignment = C if c in (1, 4, 6) else T
            cell.fill = fl(WHITE)
            cell.border = Border(bottom=Side(style="hair", color=LINE))
        if link_sheet:
            cell = ws.cell(r, 7)
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{link_sheet}'!A1", display=str(group))
            cell.font = ft(10, True, LINK, underline="single")
        self.r += 1
        return self

    def toc_row(self, ch, name, page, ids, sheet):
        ws, r = self.ws, self.r
        ws.row_dimensions[r].height = 22
        vals = [ch, name, page, ids, "Open sheet →", "", ""]
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=COLS)
        for c, v in enumerate(vals[:5], 1):
            cell = ws.cell(r, c, v)
            cell.font = ft(11, c == 1, TEXT)
            cell.alignment = L
            cell.border = Border(bottom=Side(style="hair", color=LINE))
        link = ws.cell(r, 5)
        link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{sheet}'!A1", display="Open sheet →")
        link.font = ft(11, True, LINK, underline="single")
        self.r += 1
        return self

    def _paint_range(self, r1, c1, r2, c2, bg=None):
        ws = self.ws
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                cell = ws.cell(r, c)
                if bg:
                    cell.fill = fl(bg)
                left = box_side if c == c1 else Side(style=None)
                right = box_side if c == c2 else Side(style=None)
                top = box_side if r == r1 else Side(style=None)
                bottom = box_side if r == r2 else Side(style=None)
                cell.border = Border(left=left, right=right, top=top, bottom=bottom)

    def _write_box(self, start, c1, c2, title, items, intro=None):
        """Write a titled numbered box starting at row `start`. Returns last row used."""
        ws = self.ws
        r = start
        span = c2 - c1 + 1
        chars = max(36, 16 * span)

        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        t = ws.cell(r, c1, "  " + title)
        t.font = ft(12, True, TEAL)
        t.alignment = L
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fl(HDR_BG)
        ws.row_dimensions[r].height = 22
        r += 1

        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        p = ws.cell(r, c1, ("  " + intro) if intro else "")
        p.font = ft(10, False, MUTED, italic=True)
        p.alignment = T
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = fl(WHITE)
        ws.row_dimensions[r].height = min(52, 18 + (len(intro or "") // chars) * 14) if intro else 10
        r += 1

        for i, it in enumerate(items, 1):
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            cell = ws.cell(r, c1, f"    {i}.  {it}")
            cell.font = ft(11, False, TEXT)
            cell.alignment = T
            for c in range(c1, c2 + 1):
                ws.cell(r, c).fill = fl(WHITE)
            ws.row_dimensions[r].height = min(78, 24 + (len(it) // chars) * 16)
            r += 1
        return r - 1

    def _pad_box(self, from_row, to_row, c1, c2):
        ws = self.ws
        for r in range(from_row, to_row + 1):
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            for c in range(c1, c2 + 1):
                ws.cell(r, c).fill = fl(WHITE)
            if not ws.row_dimensions[r].height:
                ws.row_dimensions[r].height = 18

    def info_box(self, title, items, intro=None):
        """Full-width numbered box with a thin black border."""
        self.space(8)
        start = self.r
        last = self._write_box(start, 1, COLS, title, items, intro)
        self._paint_range(start, 1, last, COLS)
        for c in range(1, COLS + 1):
            self.ws.cell(start, c).fill = fl(HDR_BG)
        self.r = last + 1
        self.space(8)
        return self

    def pair_boxes(self, left, right):
        """Two side-by-side numbered boxes. Each arg is (title, items) or (title, items, intro)."""
        def unpack(spec):
            if len(spec) == 2:
                return spec[0], spec[1], None
            return spec[0], spec[1], spec[2]

        lt, li, ln = unpack(left)
        rt, ri, rn = unpack(right)
        self.space(8)
        start = self.r
        last_l = self._write_box(start, 1, 5, lt, li, ln)
        last_r = self._write_box(start, 6, COLS, rt, ri, rn)
        last = max(last_l, last_r)
        if last_l < last:
            self._pad_box(last_l + 1, last, 1, 5)
        if last_r < last:
            self._pad_box(last_r + 1, last, 6, COLS)

        def ih(text, span):
            chars = max(36, 16 * span)
            return min(78, 24 + (len(text) // chars) * 16)

        self.ws.row_dimensions[start].height = 22
        intro_h = max(
            min(52, 18 + (len(ln or "") // 80) * 14) if ln else 10,
            min(52, 18 + (len(rn or "") // 64) * 14) if rn else 10,
        )
        self.ws.row_dimensions[start + 1].height = intro_h
        n_items = max(len(li), len(ri))
        for i in range(n_items):
            hl = ih(li[i], 5) if i < len(li) else 18
            hr = ih(ri[i], 4) if i < len(ri) else 18
            self.ws.row_dimensions[start + 2 + i].height = max(hl, hr)
        self._paint_range(start, 1, last, 5)
        self._paint_range(start, 6, last, COLS)
        for c in range(1, 6):
            self.ws.cell(start, c).fill = fl(HDR_BG)
        for c in range(6, COLS + 1):
            self.ws.cell(start, c).fill = fl(HDR_BG)
        self.r = last + 1
        self.space(8)
        return self

    def jump(self, label, sheet):
        """Full-width hyperlink row to another sheet."""
        self.space(4)
        cell = self._merge(1, COLS, "  " + label, size=11, bold=True, color=LINK, align=L, h=22)
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1", display=label)
        cell.font = ft(11, True, LINK, underline="single")
        self.r += 1
        return self

    def data_table(self, headers, rows, spans):
        """Document-style table. spans: list of (c1, c2) inclusive for each column."""
        ws = self.ws

        def paint_row(values, header=False, stripe=False):
            r = self.r
            chars = [max(12, 14 * (c2 - c1 + 1)) for c1, c2 in spans]
            hgt = 22
            for val, nchar in zip(values, chars):
                hgt = max(hgt, min(72, 20 + (len(str(val)) // nchar) * 14))
            ws.row_dimensions[r].height = 22 if header else hgt
            bg = NAVY if header else (HDR_BG if stripe else WHITE)
            fg = WHITE if header else TEXT
            for i, (val, (c1, c2)) in enumerate(zip(values, spans)):
                ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
                cell = ws.cell(r, c1, val)
                cell.font = ft(10 if header else 11, True if header or i == 0 else False, fg)
                cell.alignment = C if (header or i == 0) else T
                cell.fill = fl(bg)
                for c in range(c1, c2 + 1):
                    ws.cell(r, c).fill = fl(bg)
                    ws.cell(r, c).border = Border(
                        left=box_side if c == c1 else Side(style=None),
                        right=box_side if c == c2 else Side(style=None),
                        top=box_side,
                        bottom=box_side,
                    )
                cell.border = Border(left=box_side, right=box_side, top=box_side, bottom=box_side)
            self.r += 1

        paint_row(headers, header=True)
        for i, row in enumerate(rows):
            paint_row(row, stripe=(i % 2 == 0))
        self.space(8)
        return self

    def embed_figure(self, path, width_px=None):
        """Embed a document PNG at native size (or scaled to width_px). Reserves one tall row."""
        with open(path, "rb") as f:
            f.read(16)
            iw, ih = struct.unpack(">II", f.read(8))
        if width_px is None:
            width_px = iw
        height_px = int(round(width_px * ih / float(iw)))
        img = XLImage(path)
        img.width = width_px
        img.height = height_px
        self.ws.add_image(img, f"A{self.r}")
        self.ws.row_dimensions[self.r].height = max(90, height_px * 0.75)
        self.r += 1
        self.space(8)
        return self

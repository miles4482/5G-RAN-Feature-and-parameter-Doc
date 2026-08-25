#!/usr/bin/env python3
"""eRAN21.1 4G mobility — step-by-step Excel summary of three Huawei books.

Not the user SN sample. Organised as one synchronised chain, then each
feature in procedure order: Idle → Connected → Intra-RAT MLB → activation.
"""

import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.page import PageMargins
from openpyxl.drawing.image import Image as XLImage

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from mobility_figures import build_all as build_figures

OUT = "/workspace/docs/4G_LTE_Mobility_Management/4G_LTE_Mobility_Management_eRAN21.1_v3.3.xlsx"
COLS = 8

BLUE = "005596"
YELLOW = "FFFF00"
GREEN = "008000"
HDR = "DDEBF7"
GREY = "F2F2F2"
WHITE = "FFFFFF"
BLACK = "000000"
ARROW = "1F4E79"
GRID = "B4B4B4"
MAJOR_BG = "FFF2CC"

thin = Border(
    left=Side(style="thin", color=GRID),
    right=Side(style="thin", color=GRID),
    top=Side(style="thin", color=GRID),
    bottom=Side(style="thin", color=GRID),
)
C = Alignment(wrap_text=True, vertical="center", horizontal="center")
T = Alignment(wrap_text=True, vertical="top", horizontal="left")
LI = Alignment(wrap_text=True, vertical="center", horizontal="left", indent=1)

H_STEP = ["Step", "What happens", "MO / parameter", "Rule in the feature book", "If this is wrong", "Source"]
H_NOTE = ["No.", "Huawei caution", "What it means in the network", "Do / do not", "Related step", "Source"]
H_PAR = ["Order", "MO", "Parameter", "Core Setting", "Parameter Value", "Role in the procedure", "Depends on / couples with", "Source"]
H_MML = ["Order", "MO", "MML command (run in this order)", "Must already be true", "Notes", "Source"]
H_LINK = ["Open this sheet", "Book", "Issue", "This sheet answers", "Read after", "Source"]

W = [8, 20, 28, 18, 28, 32, 28, 28]


def fl(h):
    return PatternFill("solid", fgColor=h)


def ft(size=10, bold=False, color=BLACK, underline=None):
    return Font(name="Calibri", size=size, bold=bold, color=color, underline=underline)


def put(ws, r, c, v, size=10, bold=False, color=BLACK, bg=GREY, align=None, h=None):
    x = ws.cell(r, c, v)
    x.font = ft(size, bold, color)
    x.fill = fl(bg)
    x.alignment = align or T
    x.border = thin
    if h:
        ws.row_dimensions[r].height = h
    return x


def merge(ws, r, c1, c2, v, **kw):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    put(ws, r, c1, v, **kw)
    bg = kw.get("bg", GREY)
    color = kw.get("color", BLACK)
    size = kw.get("size", 10)
    bold = kw.get("bold", False)
    al = kw.get("align", T)
    for c in range(c1 + 1, c2 + 1):
        y = ws.cell(r, c)
        y.fill = fl(bg)
        y.border = thin
        y.font = ft(size, bold, color)
        y.alignment = al


def widths(ws, xs):
    for i, w in enumerate(xs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def setup(ws, footer, tab=BLUE):
    ws.sheet_view.showGridLines = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(0.35, 0.35, 0.5, 0.45)
    ws.oddHeader.left.text = "4G LTE Mobility Management  |  Huawei eRAN21.1  |  v3.3"
    ws.oddFooter.left.text = footer
    ws.oddFooter.right.text = "Page &P of &N"
    ws.sheet_properties.tabColor = tab
    ws.freeze_panes = "A3"
    ws.print_title_rows = "1:1"
    ws.sheet_format.defaultRowHeight = 18


def title(ws, r, text):
    merge(ws, r, 1, COLS, text, size=16, bold=True, color=WHITE, bg=BLUE, align=C, h=30)
    return r + 1


def spacer(ws, r, h=8):
    for c in range(1, COLS + 1):
        put(ws, r, c, "", bg=WHITE, h=h)
    return r + 1


def section(ws, r, text):
    merge(ws, r, 1, COLS, "  " + text, size=12, bold=True, color=GREEN, bg=YELLOW, align=LI, h=22)
    return r + 1


def major(ws, r, text):
    put(ws, r, 1, "NOTE", size=9, bold=True, color=WHITE, bg=GREEN, align=C, h=22)
    merge(ws, r, 2, COLS, "  " + text, size=10, bold=True, color=GREEN, bg=MAJOR_BG, align=LI, h=22)
    return r + 1


def note(ws, r, text):
    merge(ws, r, 1, COLS, "  " + text, size=9, bold=False, color=BLUE, bg=WHITE, align=LI, h=20)
    return r + 1


def auto_h(values):
    longest = 1
    for i, v in enumerate(values):
        s = "" if v is None else str(v)
        width = 16 if i == 0 else 40
        longest = max(longest, (len(s) + width - 1) // width, s.count("\n") + 1)
    return min(72, max(22, 14 + longest * 12))


def heads(ws, r, names):
    names = list(names) + [""] * COLS
    for c in range(1, COLS + 1):
        put(ws, r, c, names[c - 1], size=9, bold=True, color=BLACK, bg=HDR, align=C, h=20)
    return r + 1


def rec(ws, r, values):
    values = list(values) + [""] * COLS
    h = auto_h(values[:COLS])
    for c in range(1, COLS + 1):
        put(ws, r, c, values[c - 1], size=10, bg=GREY, align=T, h=h)
    return r + 1


def flow(ws, r, boxes, ref=""):
    """Three process boxes per row: box → box → box."""
    steps = list(boxes)
    box_spans = [(1, 2), (4, 5), (7, 8)]
    arr_cols = [3, 6]
    i = 0
    while i < len(steps):
        chunk = steps[i:i + 3]
        for c in range(1, COLS + 1):
            put(ws, r, c, "", bg=WHITE, h=28)
        for j, step in enumerate(chunk):
            c1, c2 = box_spans[j]
            merge(ws, r, c1, c2, step, size=9, bold=True, color=WHITE, bg=BLUE, align=C, h=28)
            if j < len(chunk) - 1:
                put(ws, r, arr_cols[j], "→", size=16, bold=True, color=GREEN, bg=WHITE, align=C, h=28)
        r += 1
        i += 3
        if i < len(steps):
            for c in range(1, COLS + 1):
                put(ws, r, c, "", bg=WHITE, h=12)
            merge(ws, r, 1, COLS, "↓", size=12, bold=True, color=GREEN, bg=WHITE, align=C, h=12)
            r += 1
    if ref:
        r = note(ws, r, "Document chart: " + ref)
    return r


def mml_heads(ws, r):
    put(ws, r, 1, "SN", size=9, bold=True, color=BLACK, bg=HDR, align=C, h=20)
    merge(ws, r, 2, 4, "MML", size=9, bold=True, color=BLACK, bg=HDR, align=C, h=20)
    merge(ws, r, 5, 6, "Purpose", size=9, bold=True, color=BLACK, bg=HDR, align=C, h=20)
    merge(ws, r, 7, 8, "Note", size=9, bold=True, color=BLACK, bg=HDR, align=C, h=20)
    return r + 1


def mml_row(ws, r, sn, mml, purpose, note_txt):
    h = auto_h([str(sn), mml, purpose, note_txt])
    put(ws, r, 1, sn, size=10, bold=True, bg=GREY, align=C, h=h)
    merge(ws, r, 2, 4, mml, size=9, bg=GREY, align=T, h=h)
    merge(ws, r, 5, 6, purpose, size=10, bg=GREY, align=T, h=h)
    merge(ws, r, 7, 8, note_txt, size=10, bg=GREY, align=T, h=h)
    return r + 1


def combined_mml(ws, r, rows):
    r = section(ws, r, "Combined MML Command (all Together)")
    r = note(ws, r, "All commands below are in execution order. Use Parameter Value from the table above where the feature book gives a value. Remaining <val> items must be calibrated from MR. Confirm enum names in MAE. Example A1/A2/A4 dBm in the Connected book are command examples, not design values.")
    r = mml_heads(ws, r)
    for i, (mml, purpose, note_txt) in enumerate(rows, 1):
        r = mml_row(ws, r, i, mml, purpose, note_txt)
    return r


# Feature MML in sequence. Purpose uses // as in the attached snap.
MML_IDLE = [
    ("LST CELLRESEL: LocalCellId=<x>;",
     "//Read current serving idle configuration (priority, search, serving-low, hyst) so the change has a rollback baseline.",
     "Save LST with the change record. Do not judge camping until the next SI modification period."),
    ("LST EUTRANINTERNFREQ: LocalCellId=<x>;",
     "//Read the SIB5 inter-frequency set (priority, CFG, thresholds, MeasPerformanceDemand, MlbTargetInd) before MOD.",
     "Review every DlEarfcn. A missing frequency cannot be reselected."),
    ("LST CELLALGOSWITCH: LocalCellId=<x>;",
     "//Read idle/MLB switch bits before touching InterFreqIdleMlbSwitch.",
     "Confirm Intra-RAT MLB license if idle MLB will be used."),
    ("LST RRCCONNSTATETIMER:;",
     "//Read T320ForLoadBalance before change.",
     "SPID/PCC dedicated-priority T320 is always 180 min and is not this parameter."),
    ("MOD CELLRESEL: LocalCellId=<x>, CellReselPriority=<prio>, SIntraSearchCfgInd=CFG, SNonIntraSearchCfgInd=CFG, SNonIntraSearch=10, SIntraSearch=<val_gt_10>;",
     "//Activate CFG search (required) and apply Huawei example SNonIntraSearch=10. Set SIntraSearch greater than 10. Set serving CellReselPriority so capacity/hotspot ranks above the coverage layer.",
     "Core: CfgInd=CFG. Optimized: priority ranks and SIntraSearch. Feature book does not give a numeric serving priority — design it. Idle Mode Management §5.4.1.1 ; Intra-RAT MLB §5.1.2.1"),
    ("MOD EUTRANINTERNFREQ: LocalCellId=<x>, DlEarfcn=<earfcn>, CellReselPriorityCfgInd=CFG, CellReselPriority=<prio>, MeasPerformanceDemand=NORMAL;",
     "//Publish this EARFCN in SIB5 with a priority (CFG is mandatory). NORMAL keeps the frequency visible and usable as an idle-MLB target.",
     "Repeat for every non-serving frequency that must be reselectable. UNDELIVER removes it from SIB5 and blocks idle MLB. Idle Mode Management §5.1.3.1 / §5.3.2.3"),
    ("MOD EUTRANINTERNFREQ: LocalCellId=<x>, DlEarfcn=<earfcn>, ThreshXhigh=<mr>, ThreshXlow=<mr>, QoffsetFreq=<mr>, EutranReselTime=<mr>;",
     "//Set higher-priority, lower-priority and equal-priority reselection qualification for this EARFCN. Timers must persist; camped >1 s.",
     "No universal dBm in the feature book — calibrate from MR. Too-low ThreshXhigh = premature high-band camp. Idle Mode Management Tables 5-1 to 5-4"),
    ("MOD CELLRESEL: LocalCellId=<x>, ThrshServLow=<mr>;",
     "//Allow the UE to leave a poor serving cell toward a lower-priority frequency (together with target ThreshXlow).",
     "Align with connected coverage A2/A5. Too-low serving-low keeps the UE on a dying cell. Idle Mode Management Tables 5-3/5-4"),
    ("MOD EUTRANINTERNFREQ: LocalCellId=<x>, DlEarfcn=<earfcn>, MlbTargetInd=ALLOWED;",
     "//Allow this frequency as an idle and connected MLB target. Coverage HO is still controlled by NoHoFlag/OverlapInd, not by this flag alone.",
     "Use ALLOWED_WITHOUT_IDLE_MLB or ALLOWED_WITHOUT_CONNECT_MLB to block one MLB mode. Confirm exact enum in MAE. Intra-RAT MLB pp.28, 129"),
    ("MOD CELLALGOSWITCH: LocalCellId=<x>, MlbAlgoSwitch=InterFreqIdleMlbSwitch-1;",
     "//Turn on intra-LTE idle MLB after SIB5 NORMAL, CFG priority and MlbTargetInd are already set.",
     "Core activation switch. Leave InterFreqBlindMlbSwitch-0 unless containment is proven. Confirm bit name in MAE. Idle Mode Management §5.3.2.3"),
    ("MOD RRCCONNSTATETIMER: T320ForLoadBalance=<T320>;",
     "//Set how long load-balance dedicated reselection priorities live after RRC release.",
     "Idle MLB path only. Do not expect this to change SPID/PCC T320 (180 min). Idle Mode Management §5.1.3.1"),
    ("LST CELLRESEL: LocalCellId=<x>; LST EUTRANINTERNFREQ: LocalCellId=<x>;",
     "//Verify the idle MODs were accepted, then wait SI apply delay before judging camping counters.",
     "SI is applied in the next SI modification period (or change-paging / 3 hours). SIB BER must be ≤1%. Idle Mode Management §7.1.3, §5.3.4"),
]

MML_CONN = [
    ("LST INTERFREQHOGROUP: LocalCellId=<x>, InterFreqHoGroupId=<g>;",
     "//Read A1–A5 group (coverage A2 family, A4, hyst, TTT) before any threshold MOD.",
     "Do not copy book examples A1/A2 −85/−87 dBm or A4 −103 dBm as live design. Mobility Management in Connected Mode §11.4.1.2"),
    ("LST EUTRANINTERNFREQ: LocalCellId=<x>;",
     "//Read measurement objects, FREQ_MEAS_FLAG, HO_TRG_FREQ_FORBID_MEAS_FLAG, MlbInterFreqHoEventType.",
     "Most silent A4 failures are flags/NRT, not dBm. Mobility Management in Connected Mode §4.1.4.1.2"),
    ("LST EUTRANINTERFREQNCELL: LocalCellId=<x>;",
     "//Read NRT (NoHoFlag, overlap, PCI, CIO) before enabling HO to a neighbour.",
     "Symmetric neighbour list. PERMIT_HO if coverage HO is required."),
    ("LST CELLUEMEASCONTROLCFG: LocalCellId=<x>; LST HOMEASCOMM:;",
     "//Read how many inter-frequency objects can be delivered, and SMeasure.",
     "Over-cap = random drop among equal-priority frequencies (not load balance). SMeasure can hide A4. Mobility Management in Connected Mode p.125, §4.1.5"),
    ("MOD EUTRANINTERNFREQ: LocalCellId=<x>, DlEarfcn=<earfcn>, MlbInterFreqHoEventType=A4;",
     "//Use event A4 as the MLB/FreqPri absolute target gate on this frequency (target good enough, need not beat serving). Also select FREQ_MEAS_FLAG and deselect HO_TRG_FREQ_FORBID_MEAS_FLAG in MAE for required HO targets.",
     "Core for MLB HO execution. A5 only with Intra-LTE Load Balancing for Non-cosited Cells license. Co-sited FDD typically A4. Intra-RAT MLB Table 6-3"),
    ("MOD CELLUEMEASCONTROLCFG: LocalCellId=<x>, MaxNonIntraMeasObjNum=<n>, MaxEutranFddMeasFreqNum=<n>;",
     "//Raise object capacity so every frequency this cell must measure can be delivered. n must be ≥ the number of needed inter-frequency objects.",
     "If over the cap, equal-priority objects are picked at random — that is not MLB. Mobility Management in Connected Mode Tables 4-3/4-4"),
    ("MOD HOMEASCOMM: SMeasure=<mr>;",
     "//Keep SMeasure compatible with intended A4/FreqPri so the UE does not skip inter-frequency measurement while serving is still below the A4 design.",
     "Confirm the parameter exists on this version. Too high a value silently suppresses A4. Mobility Management in Connected Mode §4.1.5"),
    ("MOD INTERFREQHOGROUP: LocalCellId=<x>, InterFreqHoGroupId=<g>, InterFreqHoA1A2Hyst=<mr>, InterFreqHoA1A2TimeToTrig=<mr>;",
     "//Set coverage A1/A2 stability from MR, then set the correct coverage A2 family (A3 vs A4/A5 vs IRAT vs blind). RSRP is the recommended quantity.",
     "Wrong A2 family = wrong HO. Book example dBm are not design values. Mobility Management in Connected Mode Tables 5-3, 5-10, 4-15"),
    ("MOD INTERFREQHOGROUP: LocalCellId=<x>, InterFreqHoGroupId=<g>, InterFreqLoadBasedHoA4ThdRsrp=<mr>, InterFreqHoA4Hyst=<mr>, InterFreqHoA4TimeToTrig=<mr>;",
     "//Set the A4 absolute target threshold better than coverage A2. This is the main MLB/FreqPri gate. Do not set A4 TTT to 5120 ms — that disables FreqPri, CQI and service-based inter-frequency HO.",
     "Calibrate from MR. 5120 ms means off, not slow. Mobility Management in Connected Mode Table 4-9 p.48 ; Table 5-22"),
    ("MOD INTRARATHOCOMM: LocalCellId=<x>, FreqPriInHoProtectionTimer=<t>, FreqPriIFHoWaitingTimer=<t>;",
     "//After an incoming unnecessary HO, hold FreqPri so the UE is not immediately bounced back. Keep FreqPri and MLB from forming a reverse pair (A→B FreqPri vs B→A MLB).",
     "Confirm exact MO names in MAE. Enable MlbBasedFreqPriHoSwitch when MLB should own heavy load. Mobility Management in Connected Mode p.300, p.303, Table 11-7"),
]

MML_MLB = [
    ("LST CELLALGOSWITCH: LocalCellId=<x>; LST CELLMLB: LocalCellId=<x>; LST CELLMLBUESEL: LocalCellId=<x>;",
     "//Read MLB master bits, trigger mode, load model, UE-pick and volume before MOD. Confirm Intra-RAT MLB license on the eNodeB.",
     "Feature book Ch.8 does not print full defaults — use the version-matched parameter reference (p.299)."),
    ("MOD EUTRANINTERNFREQ: LocalCellId=<x>, DlEarfcn=<earfcn>, MlbTargetInd=ALLOWED, MlbInterFreqHoEventType=A4;",
     "//Allow this frequency as an MLB target and execute connected load HO with event A4 (target good enough). OverlapInd and NoHoFlag=PERMIT_HO must already be valid.",
     "WITHOUT_IDLE_MLB / WITHOUT_CONNECT_MLB block one mode without blocking ordinary coverage HO. A5 needs extra license. Intra-RAT MLB pp.28, 129 ; Table 6-3"),
    ("MOD CELLMLB: LocalCellId=<x>, ActiveUeBasedLoadEvalSw=ON, SpectralEffBasedLoadEvalSw=ON, LoadTransferEnhSw=ON;",
     "//Turn on the document-recommended load model before chasing UE-number threshold: ActiveUe when MLB frequencies have different bandwidths; SpectralEff when SE differs significantly (e.g. >30%); LoadTransferEnh for multi-target math.",
     "Raw UE-count on unequal BW can reduce DL throughput. Confirm these fields/bits on this version. Intra-RAT MLB Table 5-5 ; §6.1.2.2"),
    ("MOD CELLMLB: LocalCellId=<x>, MlbTriggerMode=UE_NUMBER_ONLY, InterFreqUeTrsfType=SynchronizedUE, MlbHoCellSelectStrategy=ONLY_STRONGEST_CELL, FreqSelectStrategy=FAIRSTRATEGY;",
     "//Activate connected user-number equalisation with Huawei-recommended cell pick ONLY_STRONGEST_CELL so the UE is not immediately coverage-HO’d back. Transfer type must match the trigger mode.",
     "PRB_ONLY skips CA UEs and is burst-sensitive. FreqSelectStrategy may be PRIORITYBASED or LOADPRIORITY as designed. Intra-RAT MLB Table 6-2 ; Table 6-5 p.159"),
    ("MOD CELLMLB: LocalCellId=<x>, InterFreqMlbUeNumThd=<thd>, MlbUeNumOffset=<ofs>, MlbMaxUeNum=<n>, MlbTrigJudgePeriod=<p>, InterFreqLoadEvalPrd=<prd>;",
     "//Set enter condition N ≥ thd+offset for the whole judge period, leave when N < thd, and cap how many UEs move. Do not combine InterFreqLoadEvalPrd=5 s with MlbMaxUeNum≥40 (over-transfer).",
     "Calibrate thd from live load after ActiveUe/SE are ON. Intra-RAT MLB p.128, p.136"),
    ("MOD CELLMLBUESEL: LocalCellId=<x>;",
     "//Restrict which UEs MLB may move (UL-sync, not emergency, QCI/SPID/ARP policy). Do not select edge UEs only to reduce PRB.",
     "Confirm fields in MAE. Protect timers stop immediate re-MLB ping-pong. Intra-RAT MLB pp.130–135"),
    ("MOD CELLALGOSWITCH: LocalCellId=<x>, MlbAlgoSwitch=InterFreqMlbSwitch-1&InterFreqIdleMlbSwitch-1;",
     "//Turn on connected and idle intra-RAT MLB last, after targets, load model and thresholds exist. Idle also needs SIB5 NORMAL and T320.",
     "Core activation. Blind bit stays 0 unless containment is proven. Intra-RAT MLB Table 6-2"),
    ("MOD RRCCONNSTATETIMER: T320ForLoadBalance=<T320>;",
     "//Set idle dedicated-priority lifetime used when idle MLB releases a UE with IdleModeMobilityControlInfo.",
     "SPID/PCC T320 remains 180 min. Idle class order: NG-RAN > E-UTRAN low-load > E-UTRAN high-load > UTRAN > GERAN. Intra-RAT MLB §5.1.1.5"),
    ("LST CELLMLB: LocalCellId=<x>; LST CELLALGOSWITCH: LocalCellId=<x>;",
     "//Verify bits and thresholds, then watch Load HO / UeNumLoad, HighLoad Dur/Num, meas success, DL TP, idle DedicatedPri, and SON Inter-Frequency Handover / Idle Mode Release logs.",
     "When PRB and UE-number MLB are both on: PRB HO ≈ Load − UeNumLoad. Intra-RAT MLB Tables 6-6, 6-21 ; §§6.1.4.2, 6.5.4.2"),
]

MML_ALL = [
    ("LST CELLRESEL: LocalCellId=<x>; LST EUTRANINTERNFREQ: LocalCellId=<x>; LST EUTRANINTERFREQNCELL: LocalCellId=<x>; LST INTERFREQHOGROUP: LocalCellId=<x>, InterFreqHoGroupId=<g>; LST CELLUEMEASCONTROLCFG: LocalCellId=<x>; LST HOMEASCOMM:; LST CELLALGOSWITCH: LocalCellId=<x>; LST CELLMLB: LocalCellId=<x>; LST CELLMLBUESEL: LocalCellId=<x>; LST RRCCONNSTATETIMER:;",
     "//Dump the full mobility baseline before any MOD",
     "One sequence across Idle + Connected + MLB. Keep LST with the change record."),
] + MML_IDLE[4:9] + MML_CONN[4:] + MML_MLB[1:]


def add_fig(ws, r, key, caption, figs):
    r = note(ws, r, caption)
    path = figs[key]
    img = XLImage(path)
    scale = min(1.0, 1080 / float(img.width or 1080))
    img.width = int(img.width * scale)
    img.height = int(img.height * scale)
    img.anchor = f"A{r}"
    ws.add_image(img)
    skip = max(11, int(img.height / 18) + 2)
    for _ in range(skip):
        r = spacer(ws, r, 16)
    return r


def link_cell(cell, sheet, target="A1"):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{target}", display=str(cell.value or ""))
    cell.font = ft(10, True, "0563C1", underline="single")
    cell.alignment = T


def build_sheet(wb, name, title_text, blocks, tab=BLUE, figs=None):
    ws = wb.create_sheet(name)
    widths(ws, W)
    setup(ws, title_text, tab=tab)
    r = 1
    r = title(ws, r, title_text)
    r = spacer(ws, r, 6)
    for b in blocks:
        kind = b[0]
        if kind == "section":
            r = section(ws, r, b[1])
        elif kind == "major":
            r = major(ws, r, b[1])
        elif kind == "note":
            r = note(ws, r, b[1])
        elif kind == "flow":
            r = flow(ws, r, b[1], b[2] if len(b) > 2 else "")
        elif kind == "fig":
            r = add_fig(ws, r, b[1], b[2], figs or {})
        elif kind == "heads":
            r = heads(ws, r, b[1])
        elif kind == "row":
            rec(ws, r, b[1])
            if len(b) > 2:
                link_cell(ws.cell(r, 1), b[2])
            r += 1
        elif kind == "space":
            r = spacer(ws, r, b[1] if len(b) > 1 else 8)
        elif kind == "mml":
            r = combined_mml(ws, r, b[1])
    return ws


def overview():
    return [
        ("section", "What this file is"),
        ("note", "Step-by-step summary of three Huawei eRAN21.1 feature-parameter books. One mobility chain. Not a copy of any sample template. Values in MML examples in the books are examples, not live design."),
        ("space", 6),
        ("heads", ["Book", "Document", "Issue / date", "Owns in the chain", "Does not own", "Read as"]),
        ("row", ["Idle Mode Management", "Idle Mode Management Feature Parameter Description", "eRAN21.1 Issue 04, 2026-05-30",
                 "Where the UE camps. Next RRC setup cell. Dedicated priority at release.", "Connected HO events. Load algorithm.", "Sheet 2"]),
        ("row", ["Connected Mode", "Mobility Management in Connected Mode Feature Parameter Description", "eRAN21.1 Issue 08, 2026-06-30",
                 "Measurement and HO engine (A1–A5). Coverage first. Frequency-priority steering.", "Who to move for load. Load = N/C.", "Sheet 3"]),
        ("row", ["Intra-RAT MLB", "Intra-RAT Mobility Load Balancing Feature Parameter Description", "eRAN21.1 Issue 10, 2026-06-30",
                 "When a cell is overloaded. Which UEs. Idle dedicated-priority transfer or connected load HO.", "RF plan. Coverage A2 family design.", "Sheet 4"]),
        ("space", 8),
        ("section", "How the three books fit (one chain)"),
        ("fig", "chain", "Document chain  |  Idle Fig 5-1  →  Connected Fig 4-1  →  MLB Fig 3-1"),
        ("flow", ["Idle: camp / reselect", "RRC connect", "Coverage A2/A5 protect", "FreqPri A1/A4 steer", "MLB load HO / idle release", "Back to idle (T320)"],
         "Idle Fig 5-1  →  Connected Fig 4-1  →  MLB Fig 3-1"),
        ("major", "Idle decides the next access layer. Connected executes measurement and handover. MLB only decides who and when to move for load. Frequency-priority HO is not the MLB algorithm."),
        ("major", "Necessary coverage HO preempts load / optimisation HO. Do not use connected frequency-priority as a substitute for Intra-RAT MLB."),
        ("section", "Document charts in this file"),
        ("note", "Huawei PDF page images are not in this workspace, so each chart is rebuilt from the documented procedure and labelled with the book figure number. Open the feature sheet to see the chart next to the steps."),
        ("heads", ["Fig", "Document", "What the chart shows", "Sheet", "Use it for", "Book"]),
        ("row", ["Fig 5-1", "Idle Mode Management", "Selection / reselection sequence", "2. Idle Mode", "Camping and next RRC cell", "Issue 04"]),
        ("row", ["Fig 4-1", "Idle Mode Management", "Idle functions (select, reselect, SI, dedicated prio)", "2. Idle Mode", "Idle function map", "Issue 04"]),
        ("row", ["Tables 5-1 to 5-4", "Idle Mode Management", "Higher / equal / lower reselection decision", "2. Idle Mode", "ThreshXhigh / ThrshServLow / ranking", "Issue 04"]),
        ("row", ["Fig 4-1", "Connected Mode", "HO procedure A1–A5", "3. Connected Mode", "Measurement and HO engine", "Issue 08"]),
        ("row", ["A1–A5", "Connected Mode", "Event meaning", "3. Connected Mode", "Coverage vs MLB/FreqPri gate", "Issue 08"]),
        ("row", ["Fig 11-1 / 11-2", "Connected Mode", "Frequency-priority HO (not MLB)", "3. Connected Mode", "High-band steering", "Issue 08"]),
        ("row", ["Fig 3-1", "Intra-RAT MLB", "Load balancing procedure", "4. Intra-RAT MLB", "Who/when to move", "Issue 10"]),
        ("row", ["Fig 4-1", "Intra-RAT MLB", "Equalisation vs offload", "4. Intra-RAT MLB", "Need peer load or not", "Issue 10"]),
        ("row", ["Figs 4-4 / 4-5", "Intra-RAT MLB", "Idle transfer vs connected transfer", "4. Intra-RAT MLB", "T320 vs A4/A5 HO", "Issue 10"]),
        ("space", 8),
        ("section", "Open the sheets in this order. Each sheet has the same layout: overview, document chart, procedure table, Combined MML at the bottom."),
        ("heads", H_LINK),
        ("row", ["1. End-to-end chain", "All three", "—", "The full sequence from power-on to the next idle camp", "Start here", "This file"], "1. End-to-end chain"),
        ("row", ["2. Idle Mode", "Idle Mode Management", "Issue 04", "Selection, SIB, reselection, dedicated priority", "After the chain", "Idle book"], "2. Idle Mode"),
        ("row", ["3. Connected Mode", "Connected Mode", "Issue 08", "A1–A5, coverage HO, frequency-priority HO", "After Idle", "Connected book"], "3. Connected Mode"),
        ("row", ["4. Intra-RAT MLB", "Intra-RAT MLB", "Issue 10", "Load model, trigger, UE pick, A4/A5 or idle release", "After Connected", "MLB book"], "4. Intra-RAT MLB"),
        ("row", ["5. Activation order", "All three", "—", "One MML sequence across Idle + Connected + MLB", "Last", "MAE / parameter reference"], "5. Activation order"),
        ("space", 8),
        ("section", "What is not in this file"),
        ("row", ["1", "No AI / daily-KPI / change-request agent", "—", "Feature-book summary only", "—", "—"]),
        ("row", ["2", "No operator-specific band priority table", "—", "Huawei says capacity/hotspot above coverage layer. Exact priority numbers are a design, not a book default.", "Idle + MLB", "MLB §5.1.2.1"]),
        ("row", ["3", "No full default/range list", "—", "MLB Ch.8 points to the version-matched parameter reference.", "Activation sheet", "MLB p.299"]),
        ("row", ["4", "MML uses LocalCellId=<x> and DlEarfcn=<earfcn>", "—", "Confirm enum names and syntax on the NE in MAE-Access.", "Last section of every sheet", "MAE"]),
        ("space", 8),
        ("mml", MML_ALL),
    ]


def chain():
    return [
        ("section", "Read this as one procedure"),
        ("note", "Each row is the next thing that happens to a UE. The Book column tells you which feature document owns that step. Open sheets 2–4 for the detail of that step."),
        ("fig", "chain", "End-to-end chain  |  Idle Fig 5-1  →  Connected Fig 4-1  →  MLB Fig 3-1"),
        ("flow", ["Power-on / idle", "Camped cell → RRC", "Protect coverage", "Steer to high band", "Balance load", "Release → next idle"],
         "Synchronised chain of the three books"),
        ("space", 6),
        ("heads", ["Step", "What happens", "Which book", "Trigger", "What the eNodeB / UE does", "Go to"]),
        ("row", ["1", "PLMN selection then cell selection", "Idle", "Power-on, recovery, or no stored cell", "Criterion S: Srxlev > 0 (and Squal if QQualMin is set). This is a floor, not a load knob.", "Sheet 2 step 1–3"]),
        ("row", ["2", "UE camps and reads system information", "Idle", "Suitable cell found", "SIB1 (80 ms). SIB3 = intra + serving priority. SIB5 = inter-frequency priority, thresholds, neighbour list.", "Sheet 2 step 4"]),
        ("row", ["3", "Idle measurement and reselection", "Idle", "Always for higher-priority freqs; equal/lower after search threshold", "Higher: ThreshXhigh. Equal: ranking Rn/Rs. Lower: ThrshServLow and ThreshXlow. Timers must persist.", "Sheet 2 step 5–9"]),
        ("row", ["4", "RRC setup on the camped cell", "Idle → Connected", "Service request", "Dedicated idle priorities are discarded when the UE enters connected. The camped cell is the serving cell of this session.", "Idle §5.1.3.1"]),
        ("row", ["5", "Classify the HO function", "Connected", "Coverage / service / quality / FreqPri / MLB", "Necessary coverage HO has higher priority than unnecessary load or optimisation HO.", "Connected Table 4-5"]),
        ("row", ["6", "Deliver measurement (or blind only if containment is known)", "Connected", "HO function started", "FREQ_MEAS_FLAG, not forbid-meas, object-count limits. Equal-priority objects may be picked at random if over cap — that is not load balance.", "Connected Fig 4-1"]),
        ("row", ["7", "Coverage protection", "Connected", "Serving crosses A2", "A1 stops coverage meas. A3 relative / A4 target good enough / A5 serving poor AND target good. RSRP is the recommended quantity.", "Sheet 3"]),
        ("row", ["8", "Frequency-priority HO (optional)", "Connected Ch.11", "Serving is good enough (A1) and a high-priority frequency is usable (A4)", "Puts service on high band and keeps low band for coverage. A frequency that is a FreqPri target must not have a reverse MLB-target pair.", "Fig 11-1 / 11-2"]),
        ("row", ["9", "MLB evaluates load", "MLB", "Periodic eval", "User-number load = N/C. Difference = (Load_s − Load_t) / Load_s. Use ActiveUe + SpectralEff when bandwidth or SE differ.", "Sheet 4"]),
        ("row", ["10", "MLB trigger and target admit", "MLB", "N ≥ thd+offset for the judge period (or PRB path)", "Target must pass NoHoFlag, overlap, MLB target indication, HO success, Low/Med HW+transport.", "MLB pp.27–29"]),
        ("row", ["11", "MLB picks UEs and volume", "MLB", "After admit", "ONLY_STRONGEST_CELL is recommended. Cap volume (do not use 5 s eval with MlbMaxUeNum ≥ 40).", "Table 6-5, p.136"]),
        ("row", ["12", "Connected load HO executes", "MLB + Connected", "UE selected", "FDD: measurement-based HO, event A4 (or A5 if licensed non-cosited). Cause: Reduce Load in Serving Cell.", "MLB §6.1.1.5.2"]),
        ("row", ["13", "Idle load transfer at release", "MLB + Idle", "Idle MLB trigger", "RRC release with dedicated reselection priority and T320. Class order: NG-RAN > E-UTRAN low-load > E-UTRAN high-load > UTRAN > GERAN.", "MLB §5.1.1.5"]),
        ("row", ["14", "Back to idle / next session", "Idle", "Release, T320 expiry, or next RRC", "Dedicated priorities last until T320, next RRC, or PLMN select. SI change is not instant — wait the next SI modification period.", "Idle §7.1.3"]),
        ("space", 8),
        ("section", "Priority between these steps"),
        ("row", ["A", "Coverage always wins", "Connected", "A2/A5 running", "Do not let MLB or FreqPri pull a UE off a dying serving cell.", "Table 4-5"]),
        ("row", ["B", "FreqPri and MLB must not fight", "Connected + MLB", "Same frequency pair", "No reverse MLB target on a FreqPri pair. MlbBasedFreqPriHoSwitch lets MLB own heavy load.", "Connected p.303"]),
        ("row", ["C", "Idle MLB and connected user-number MLB", "Idle + MLB", "Both ON", "Do not combine fixed-proportion idle MLB with user-number connected MLB (ping-pong).", "MLB §5.4.2.2"]),
        ("row", ["D", "Admission", "Connected", "Unnecessary / offload HO", "Target must admit ALL QCIs. Prep fail is often admission, not RF.", "Tables 4-16 / 4-17"]),
        ("space", 8),
        ("mml", MML_ALL),
    ]


def idle():
    return [
        ("section", "Document charts"),
        ("fig", "idle_41", "Idle Mode Management  Fig 4-1  — idle functions"),
        ("fig", "idle_51", "Idle Mode Management  Fig 5-1  — selection / reselection sequence"),
        ("fig", "idle_reselect", "Idle Mode Management  Tables 5-1 to 5-4  — reselection decision"),
        ("section", "Procedure chart"),
        ("flow", ["PLMN select", "Criterion S", "Camp + read SI", "Measure", "Reselect", "RRC on camped cell"],
         "Idle Mode Management Fig 4-1 and Fig 5-1"),
        ("space", 6),
        ("section", "Step by step"),
        ("heads", H_STEP),
        ("row", ["1", "PLMN selection", "NAS / PLMN", "Happens before cell selection. SIB reselection priority does not control this.", "Wrong PLMN looks like a reselection fault.", "Fig 5-1"]),
        ("row", ["2", "Initial / stored / release-directed cell selection", "CELLSEL ; RRCConnectionRelease redirect", "Last camped, stored info, or strongest cell per band. SIB priority does not fully control first camp after power-up.", "Do not expect SIB priority to fix first access after reset.", "§5.1.2"]),
        ("row", ["3", "Criterion S (suitable cell)", "QRxLevMin, QQualMin, UePowerMax / PMax", "Srxlev = Qrxlevmeas − (Qrxlevmin + offset) − Pcompensation > 0. Squal > 0 only if QQualMin is set.", "Too-high floor = access failure. This is not a load-balance knob.", "§5.1.2"]),
        ("row", ["4", "Camp and read SI", "SIB1 / SIB3 / SIB4 / SIB5", "SIB3: intra + serving common priority. SIB5: inter-freq priority, thresholds, offset, neighbour list. Max 16 non-serving E-UTRAN frequencies.", "Missing SIB5 frequency = UE cannot reselect there.", "Table 7-1"]),
        ("row", ["5", "Intra-frequency measurement", "SIntraSearch / SIntraSearchQ", "UE measures intra unless serving is above the search threshold. Huawei: set CfgInd = CFG.", "Always-on intra meas if not CFG.", "§5.1.3.3"]),
        ("row", ["6", "Inter-frequency measurement", "SNonIntraSearch ; CellReselPriority", "Higher-priority frequency: always measured. Equal/lower: only after search threshold. Search threshold does not stop high-priority search.", "Huawei example SNonIntraSearch = 10; SIntraSearch > SNonIntraSearch.", "§5.1.3.3 ; §5.4.1.1"]),
        ("row", ["7", "Reselection to a higher-priority frequency", "ThreshXhigh / ThreshXhighQ / EutranReselTime", "Camped > 1 s AND target S > high threshold for the timer.", "Too-low ThreshXhigh = premature high-band camp.", "Tables 5-1 / 5-2"]),
        ("row", ["8", "Reselection to equal-priority frequency", "Qhyst, QoffsetFreq, CellQoffset", "Rn = Qmeas,n − Qoffset ; Rs = Qmeas,s + Qhyst ; Rn > Rs for the timer. Ranking is RF, not load.", "Do not use static offset as hourly load control.", "§5.1.3.4"]),
        ("row", ["9", "Reselection to a lower-priority frequency", "ThrshServLow / ThreshXlow", "No higher target AND serving S < serving-low AND target S > X-low, for the timer.", "Too-low serving-low = UE stays on a dying serving cell.", "Tables 5-3 / 5-4"]),
        ("row", ["10", "Dedicated priority at RRC release", "IdleModeMobilityControlInfo ; T320", "Replaces common SIB priorities until T320, next RRC, or PLMN select. Missing serving freq in the dedicated list is treated as lowest.", "Idle MLB / SPID / PCC use this path. SPID/PCC T320 is always 180 min.", "§5.1.3.1"]),
        ("row", ["11", "Next RRC setup", "Camped cell", "UE establishes RRC on the cell where it is then camped. Idle distribution = next access layer.", "A long connected session cannot be moved by idle MLB.", "Fig 5-1"]),
        ("space", 8),
        ("section", "Huawei cautions (read with the steps above)"),
        ("heads", H_NOTE),
        ("row", ["1", "Priority is per frequency, not per cell", "No SIB5 priority ⇒ no reselection to that frequency. Max 16 non-serving E-UTRAN frequencies.", "Set CellReselPriorityCfgInd = CFG on every frequency the UE must be able to reselect.", "Step 4, 7–9", "§5.1.3.1"]),
        ("row", ["2", "Capacity / hotspot above coverage layer", "Huawei recommendation for broadcast reselection priority.", "Do not publish operator-specific numbers as book defaults. Design the ranks; keep coverage layer under capacity layers.", "Step 7", "MLB §5.1.2.1"]),
        ("row", ["3", "UNDELIVER blocks idle MLB", "MeasPerformanceDemand = UNDELIVER removes the frequency from SIB5. It cannot be an idle-MLB target.", "Keep NORMAL on frequencies idle MLB must use.", "Step 4, 10", "§5.3.2.3"]),
        ("row", ["4", "SI is delayed", "Applied in the next SI modification period; otherwise after change-paging or 3 hours. SIB BER must be ≤ 1%.", "Do not judge camping in the same 15-min as the change.", "Step 4", "§7.1.3, §5.3.4"]),
        ("row", ["5", "Do not use the optimisation-table sentence on QRxLevMinOffset", "It conflicts with the Criterion-S formula.", "Treat Criterion S as written in §5.1.2.", "Step 3", "§5.1.2 vs §5.4.1.1.3"]),
        ("row", ["6", "Adaptive-proportion idle balancing is not recommended without Huawei support", "Fixed-proportion idle + user-number connected MLB causes ping-pong.", "Keep that pair off together.", "Step 10", "MLB §5.5.2.1 ; §5.4.2.2"]),
        ("row", ["7", "Neighbour list cap 16", "SIB listed neighbours are truncated.", "Truncation looks like a threshold problem.", "Step 4", "§5.1.3.4"]),
        ("space", 8),
        ("section", "What idle gives / where it stops"),
        ("heads", ["Type", "Item", "Document statement", "Condition", "User impact", "Source"]),
        ("row", ["Gives", "Low signalling vs connected HO", "Idle transfer avoids gap-assisted inter-frequency measurement and connected HO.", "Dedicated priority / idle MLB", "Lower immediate gap loss", "MLB Figs 4-4 / 4-5"]),
        ("row", ["Gives", "Sets next access cell", "Next RRC starts on the camped cell.", "After reselection or dedicated prio", "Long-term layer mix", "Fig 5-1"]),
        ("row", ["Stops", "Not real-time", "Dedicated priorities discarded at connected / T320 / PLMN select.", "Long RRC session", "Cannot rebalance an already-connected UE", "§5.1.3.1"]),
        ("row", ["Stops", "Highest common priority collects most idle UEs", "Equal-priority still ranks by RF, not load.", "No idle MLB", "Need idle MLB for dynamic load", "§5.1.3.5"]),
        ("space", 8),
        ("section", "Parameters in the same order as the procedure"),
        ("note", "One parameter per row. Core parameter = required to activate the function. Basic and Optimized = used to tune performance after activation. Parameter Value is from the Huawei feature book; if the book gives no default, calibrate from MR."),
        ("heads", H_PAR),
        ("row", ["1", "CELLSEL", "QRxLevMin", "Basic and Optimized",
                 "No universal dBm in the book. Set so Srxlev>0 on a suitable cell.",
                 "Selection RSRP floor (Criterion S)", "Not a load-balance knob", "Idle Mode Management §5.1.2"]),
        ("row", ["2", "CELLSEL", "QQualMin", "Basic and Optimized",
                 "0 or absent = RSRP-only. If configured, Squal>0 is also required.",
                 "Selection RSRQ floor (Criterion S)", "RSRQ moves with load and can oscillate", "Idle Mode Management §5.1.2"]),
        ("row", ["3", "CELLSEL", "QRxLevMinOffset", "Basic and Optimized",
                 "Used as (Qrxlevmin + offset) in Srxlev. Do not use the optimisation-table sentence (it conflicts with the Criterion-S formula).",
                 "Selection offset", "Follow §5.1.2, not §5.4.1.1.3", "Idle Mode Management §5.1.2"]),
        ("row", ["4", "CELLSEL", "UePowerMax", "Basic and Optimized",
                 "Used in Pcompensation in Criterion S",
                 "UE max TX for suitability", "Pair with PMax on CELLRESEL", "Idle Mode Management §5.1.2"]),
        ("row", ["5", "CELLRESEL", "QRxLevMin", "Basic and Optimized",
                 "Same philosophy as selection. No universal dBm in the book.",
                 "Serving / intra reselection suitability", "Not a load knob", "Idle Mode Management §5.1.3.4"]),
        ("row", ["6", "CELLRESEL", "QQualMin", "Basic and Optimized",
                 "0 or absent = RSRP-only reselection suitability",
                 "Serving RSRQ floor", "RSRQ can oscillate with load", "Idle Mode Management §5.1.3.4"]),
        ("row", ["7", "CELLRESEL", "PMax", "Basic and Optimized",
                 "Serving PMax for suitability",
                 "Reselection Pcompensation", "Pair with UePowerMax", "Idle Mode Management §5.1.3.4"]),
        ("row", ["8", "CELLRESEL", "CellReselPriority", "Basic and Optimized",
                 "Larger = higher. Huawei: capacity/hotspot above coverage layer. No numeric default in the book.",
                 "SIB3 serving common priority", "Frequency-level, not per-cell. Max 16 non-serving E-UTRAN frequencies.", "Idle Mode Management §5.1.3.1 ; Intra-RAT MLB §5.1.2.1"]),
        ("row", ["9", "EUTRANINTERNFREQ", "CellReselPriorityCfgInd", "Core parameter",
                 "CFG",
                 "Without CFG the UE does not reselect to that frequency", "Must be CFG on every frequency that must be reselectable", "Idle Mode Management §5.1.3.1"]),
        ("row", ["10", "EUTRANINTERNFREQ", "CellReselPriority", "Basic and Optimized",
                 "Same hierarchy as serving design. No numeric default in the book.",
                 "SIB5 target priority", "Requires CfgInd=CFG first", "Idle Mode Management §5.1.3.1"]),
        ("row", ["11", "CELLRESEL", "SIntraSearchCfgInd", "Core parameter",
                 "CFG",
                 "Configure intra search instead of always-on intra measurement", "Pair with SIntraSearch", "Idle Mode Management §5.4.1.1"]),
        ("row", ["12", "CELLRESEL", "SIntraSearch", "Basic and Optimized",
                 "Must be greater than SNonIntraSearch",
                 "Skip intra meas when serving is very good", "CfgInd=CFG", "Idle Mode Management §5.4.1.1"]),
        ("row", ["13", "CELLRESEL", "SNonIntraSearchCfgInd", "Core parameter",
                 "CFG",
                 "Start equal/lower inter-frequency search", "Does not stop higher-priority measurement", "Idle Mode Management §5.4.1.1"]),
        ("row", ["14", "CELLRESEL", "SNonIntraSearch", "Basic and Optimized",
                 "Huawei example value 10",
                 "Equal/lower inter-frequency search start", "Example, still confirm on the NE. SIntraSearch > this value.", "Idle Mode Management §5.4.1.1 ; Intra-RAT MLB §5.1.2.1"]),
        ("row", ["15", "EUTRANINTERNFREQ", "ThreshXhigh", "Basic and Optimized",
                 "Calibrate from MR. No universal dBm in the book.",
                 "Higher-priority target qualification", "Camped >1 s and EutranReselTime must persist", "Idle Mode Management Tables 5-1/5-2"]),
        ("row", ["16", "EUTRANINTERNFREQ", "ThreshXhighQ", "Basic and Optimized",
                 "Only if RSRQ-based reselection is used. Calibrate from MR.",
                 "Higher-priority RSRQ qualification", "RSRQ can oscillate with load", "Idle Mode Management Tables 5-1/5-2"]),
        ("row", ["17", "EUTRANINTERNFREQ", "ThreshXlow", "Basic and Optimized",
                 "Calibrate from MR. No universal dBm in the book.",
                 "Lower-priority target qualification", "Used with ThrshServLow", "Idle Mode Management Tables 5-3/5-4"]),
        ("row", ["18", "CELLRESEL", "ThrshServLow", "Basic and Optimized",
                 "Calibrate from MR. Align with connected coverage A2/A5 philosophy.",
                 "Permission to leave serving toward lower priority", "Too low = UE stays on a dying serving cell", "Idle Mode Management Tables 5-3/5-4"]),
        ("row", ["19", "CELLRESEL", "Qhyst", "Basic and Optimized",
                 "Calibrate from MR. Adds stickiness to serving cell in equal-priority ranking.",
                 "Serving rank Rs = Qmeas,s + Qhyst", "Small timer + small hyst = ping-pong", "Idle Mode Management §5.1.3.4"]),
        ("row", ["20", "EUTRANINTERNFREQ", "QoffsetFreq", "Basic and Optimized",
                 "Calibrate from MR. Positive offset makes reselection to the target harder.",
                 "Equal-priority frequency offset", "Do not use static offset as hourly load control", "Idle Mode Management §5.1.3.4"]),
        ("row", ["21", "EUTRANINTERNFREQ", "EutranReselTime", "Basic and Optimized",
                 "Must persist. Calibrate from MR.",
                 "Inter-frequency reselection timer", "Too small = ping-pong", "Idle Mode Management §5.1.3.4"]),
        ("row", ["22", "CELLRESEL", "TReselEutran", "Basic and Optimized",
                 "Must persist. Calibrate from MR.",
                 "Intra-frequency reselection timer", "Too small = ping-pong", "Idle Mode Management §5.1.3.4"]),
        ("row", ["23", "EUTRANINTERFREQNCELL", "CellQoffset", "Basic and Optimized",
                 "Per-neighbour idle offset. Max 16 listed neighbours per frequency.",
                 "Neighbour rank in equal-priority", "Truncation looks like a threshold problem", "Idle Mode Management §5.1.3.4"]),
        ("row", ["24", "EUTRANINTERNFREQ", "MeasPerformanceDemand", "Core parameter",
                 "NORMAL for main capacity and intensive-coverage frequencies",
                 "SIB5 visibility", "UNDELIVER removes the frequency from SIB5 and it cannot be an idle-MLB target", "Idle Mode Management §5.1.3.3 ; §5.3.2.3"]),
        ("row", ["25", "EUTRANINTERNFREQ", "MlbTargetInd", "Core parameter",
                 "ALLOWED (or ALLOWED_WITHOUT_IDLE_MLB / ALLOWED_WITHOUT_CONNECT_MLB)",
                 "Whether this frequency may be an idle and/or connected MLB target", "Coverage NoHoFlag remains separate. Confirm enum in MAE.", "Intra-RAT MLB pp.28, 129"]),
        ("row", ["26", "CELLALGOSWITCH", "InterFreqIdleMlbSwitch", "Core parameter",
                 "ON (InterFreqIdleMlbSwitch-1)",
                 "Enables intra-LTE idle load equalisation", "Requires license, InterFreqMlbSwitch as documented, SIB5 NORMAL, idle-allowed MlbTargetInd", "Idle Mode Management §5.3.2.3"]),
        ("row", ["27", "CELLALGOSWITCH", "InterFreqMlbSwitch", "Core parameter",
                 "ON if connected MLB is also used",
                 "Master intra-RAT MLB bit used together with idle MLB", "License first", "Intra-RAT MLB Table 6-2"]),
        ("row", ["28", "RRCCONNSTATETIMER", "T320ForLoadBalance", "Basic and Optimized",
                 "Set the dedicated-priority lifetime used by idle MLB. SPID/PCC path is always 180 min.",
                 "How long load-balance dedicated priorities live", "Idle MLB ON. Discarded at next RRC / PLMN select / T320 expiry.", "Idle Mode Management §5.1.3.1"]),
        ("row", ["29", "EnhancedMlbAlgoSwitch", "DediPrioManageOnLowLoadSw", "Basic and Optimized",
                 "ON to keep released UEs off higher-load frequencies (verify license)",
                 "Hold low-load dedicated priority", "T320 must be set. Intra-RAT MLB Table 5-10", "Intra-RAT MLB §5.3.1"]),
        ("row", ["30", "GlobalProcSwitch", "CellReselectionOptSwitch", "Basic and Optimized",
                 "ON recommended with idle equalisation (preferential LTE frequency delivery)",
                 "Dedicated-priority delivery helper", "Idle MLB path", "Idle Mode Management §5.3.2.3"]),
        ("space", 8),
        ("section", "How Idle locks to the other two books"),
        ("heads", H_STEP),
        ("row", ["→ Connected", "Dedicated idle prio discarded at RRC connect", "Align ThreshXhigh with A1-escape and ThrshServLow with A5-return", "Do not change idle and coverage A2 on the same day without a design", "Sheet 3", "§5.1.3.1"]),
        ("row", ["→ MLB", "Idle transfer method = dedicated prio + T320", "InterFreqIdleMlbSwitch + idle-allowed MlbTargetInd + SIB5 NORMAL", "Class order includes UTRAN/GERAN — not a dump tool", "Sheet 4", "MLB §5.1.1.5"]),
        ("space", 8),
        ("section", "Idle activation sequence"),
        ("heads", H_MML),
        ("row", ["0", "LST", "Dump CELLRESEL, EUTRANINTERNFREQ, CELLALGOSWITCH, RRCCONNSTATETIMER", "Read-only", "Keep the LST with the change record", "MAE"]),
        ("row", ["1", "CELLRESEL", "Serving priority + CFG search; SIntraSearch > SNonIntraSearch", "SIB3 plan ready", "Huawei example SNonIntraSearch=10. Capacity/hotspot above coverage.", "§5.4.1.1"]),
        ("row", ["2", "EUTRANINTERNFREQ", "Every needed EARFCN: CfgInd=CFG, priority, MeasPerformanceDemand=NORMAL", "Frequency exists in SIB5 plan", "Repeat per non-serving frequency", "§5.1.3.1"]),
        ("row", ["3", "EUTRANINTERNFREQ", "ThreshXhigh / ThreshXlow / QoffsetFreq / EutranReselTime from MR", "Priority ranks already set", "No universal dBm in the feature book", "Tables 5-1 to 5-4"]),
        ("row", ["4", "CELLRESEL", "ThrshServLow aligned with coverage A2/A5 philosophy", "Connected coverage design known", "Do not set so low that a dying cell never yields", "Tables 5-3 / 5-4"]),
        ("row", ["5", "EUTRANINTERNFREQ", "MlbTargetInd as designed (idle allow or forbid)", "Coverage NoHoFlag remains PERMIT if coverage HO is required", "WITHOUT_IDLE_MLB blocks only idle MLB, not coverage HO", "pp.28, 129"]),
        ("row", ["6", "CELLALGOSWITCH then T320", "Idle MLB bit after 1–5. Then T320ForLoadBalance.", "License present; Blind bit stays off unless designed", "Wait next SI modification period before judging camp", "§7.1.3"]),
        ("space", 8),
        ("mml", MML_IDLE),
    ]


def connected():
    return [
        ("section", "Document charts"),
        ("fig", "conn_41", "Mobility Management in Connected Mode  Fig 4-1  — HO procedure"),
        ("fig", "conn_events", "Connected Mode  A1–A5 events  |  RSRP recommended  |  Tables 4-8, 5-16, 5-22, 5-18"),
        ("fig", "conn_freqpri", "Connected Mode  Fig 11-1 / 11-2  — frequency-priority HO (not MLB)"),
        ("section", "Procedure chart"),
        ("flow", ["Start HO function", "Meas or blind", "Deliver meas config", "UE reports A1–A5", "Pick target + admit", "Execute HO"],
         "Connected Mode Fig 4-1 §§4.1.1–4.1.8. Full MLB algorithm is in the MLB book, not here."),
        ("space", 6),
        ("section", "Step by step"),
        ("heads", H_STEP),
        ("row", ["1", "Start the HO function and classify it", "Coverage / FreqPri / MLB / service / quality", "Necessary coverage vs unnecessary offload vs unnecessary optimisation.", "Wrong class = load HO stealing a coverage rescue.", "pp.21–28 ; Table 4-5"]),
        ("row", ["2", "Choose measurement-based or blind", "Fig 4-2 ; BlindHoPriority", "Blind only when mobility is immediate and the neighbour contains the serving coverage.", "Higher access-failure risk.", "pp.31–32"]),
        ("row", ["3", "Deliver measurement configuration", "EUTRANINTERNFREQ + report config", "FREQ_MEAS_FLAG selected; HO_TRG_FREQ_FORBID deselected for required targets; object-count limits.", "Over cap: equal-priority frequencies may drop at random — not load balance.", "Tables 4-1 to 4-4, p.125"]),
        ("row", ["4", "Event A1 — serving becomes good", "A1: Ms − Hys > Thresh for TTT", "Stops coverage measurements. Can start A1-based FreqPri.", "Keep A1/A2 hysteresis consistent.", "Table 4-8"]),
        ("row", ["5", "Event A2 — serving becomes poor", "A2: Ms + Hys < Thresh for TTT", "Starts inter-frequency coverage measurement. Separate A2 families for A3, A4/A5, IRAT, blind.", "Wrong A2 family = wrong HO.", "Tables 5-1, 5-3, 5-10"]),
        ("row", ["6", "Event A3 — neighbour relatively better", "Mn+Ofn+Ocn−Hys > Ms+Ofs+Ocs+A3Offset", "Ofn/Ofs frequency offsets; Ocn/Ocs CIO.", "Large CIO can mask RF overshoot.", "Table 5-16"]),
        ("row", ["7", "Event A4 — neighbour absolutely good", "Mn+Ofn+Ocn−Hys > Thresh", "Used by MLB and FreqPri: target need only be good enough, not better than serving.", "A4 must be better than coverage A2 or the UE ping-pongs.", "Table 5-22 ; Table 11-5"]),
        ("row", ["8", "Event A5 — serving poor AND neighbour good", "Ms+Hys < Th1 AND Mn+Ofn+Ocn−Hys > Th2", "Strongest coverage-protection semantics. MLB A5 needs extra non-cosited license.", "Use A4 for co-sited FDD MLB unless that license exists.", "Tables 5-18 / 5-19"]),
        ("row", ["9", "Pick target and admit", "Best filtered target", "Necessary HO: any QCI. Unnecessary / offload: ALL QCIs. Inter-eNB unnecessary HO does not immediately try the next target after admit fail.", "Prep fail is often admission, not RF.", "Tables 4-16 / 4-17"]),
        ("row", ["10", "Execute, then punish / retry on fail", "Res / Opt / NonRes punish timers", "Do not treat every prep fail as an RF-threshold problem.", "Fix X2/admission before tightening A4.", "pp.61–65"]),
        ("space", 8),
        ("section", "Huawei cautions"),
        ("heads", H_NOTE),
        ("row", ["1", "Use RSRP as the general trigger quantity", "RSRQ moves with scheduler load and can oscillate.", "Keep RSRP unless there is a documented RSRQ design.", "Steps 4–8", "Table 4-15"]),
        ("row", ["2", "A4 must be better than the relevant coverage A2", "Otherwise coverage and load HO fight.", "Set A4 after coverage A2, from MR, not from MML examples.", "Step 7", "Table 5-22 ; Table 11-8"]),
        ("row", ["3", "A4 TTT = 5120 ms disables FreqPri, CQI and service-based inter-frequency HO", "Treat 5120 ms as off, not as slow.", "Use a real TTT if FreqPri or MLB A4 is required.", "Step 7", "Table 4-9 p.48"]),
        ("row", ["4", "Example A1/A2 −85/−87 dBm and A4 −103 dBm are command examples, not design values", "Copying them live is a design error.", "Calibrate from MR.", "Steps 4–8", "§11.4.1.2"]),
        ("row", ["5", "Equal-priority measurement objects may be selected randomly", "That is not load balance.", "Raise MaxNonIntraMeasObjNum / MaxEutranFddMeasFreqNum to cover every needed frequency.", "Step 3", "§5.3.1.2 p.125"]),
        ("row", ["6", "FreqPri target A→B must not have reverse MLB target B→A", "Ping-pong warning.", "Remove the reverse pair. Use MlbBasedFreqPriHoSwitch when MLB should own heavy load.", "Step 1, 8", "p.303 ; Table 11-7"]),
        ("row", ["7", "SMeasure can silently suppress A4", "UE may skip inter-freq meas while serving RSRP is above SMeasure.", "Check HOMEASCOMM if A4 never arrives.", "Step 3", "§4.1.5"]),
        ("space", 8),
        ("section", "What connected mobility gives / where it stops"),
        ("heads", ["Type", "Item", "Document statement", "Condition", "User impact", "Source"]),
        ("row", ["Gives", "Coverage rescue", "A2 → A3/A4/A5 avoids drop", "Necessary HO priority", "Indoor / edge", "Table 4-5"]),
        ("row", ["Gives", "A4 offload gate", "Target need only be good enough", "MLB / FreqPri", "Can leave a still-usable source", "Table 5-22"]),
        ("row", ["Gives", "FreqPri + A1", "Service on high band; low band kept for coverage", "Fig 11-1 / 11-2", "Coverage-layer protection intent", "pp.299–300"]),
        ("row", ["Stops", "Not a load algorithm", "Equal static priority ≠ instantaneous load share", "Need MLB book", "Random object pick among equal prio", "p.125"]),
        ("row", ["Stops", "Measurement gaps steal DL TTIs", "Gap pattern cost", "Old UE / VoLTE", "Prefer A1/A2 gated meas", "Fig 4-10"]),
        ("space", 8),
        ("section", "Parameters in the same order as the procedure"),
        ("note", "One parameter per row. Core parameter = required to activate the function. Basic and Optimized = used to tune performance after activation. Parameter Value is from the Huawei feature book; if the book gives no default, calibrate from MR."),
        ("heads", H_PAR),
        ("row", ["1", "EUTRANINTERNFREQ", "FREQ_MEAS_FLAG", "Core parameter",
                 "Selected for frequencies that must be measured",
                 "Whether the UE is configured to measure this frequency", "Silent no-HO if this is off", "Mobility Management in Connected Mode §4.1.4.1.2"]),
        ("row", ["2", "EUTRANINTERNFREQ", "HO_TRG_FREQ_FORBID_MEAS_FLAG", "Core parameter",
                 "Deselected for required HO targets",
                 "Forbids the frequency as a HO target if selected", "Audit flags before any dBm change", "Mobility Management in Connected Mode §4.1.4.1.2"]),
        ("row", ["3", "CELLUEMEASCONTROLCFG", "MaxNonIntraMeasObjNum", "Core parameter",
                 "≥ number of inter-frequency objects this cell must measure",
                 "How many non-intra objects can be delivered", "Over limit: equal-priority frequencies drop at random (not load balance)", "Mobility Management in Connected Mode Tables 4-3/4-4 ; p.125"]),
        ("row", ["4", "CELLUEMEASCONTROLCFG", "MaxEutranFddMeasFreqNum", "Core parameter",
                 "≥ number of FDD E-UTRAN frequencies this cell must measure",
                 "FDD meas-frequency cap", "Same random-drop risk", "Mobility Management in Connected Mode Tables 4-3/4-4"]),
        ("row", ["5", "HOMEASCOMM", "SMeasure", "Basic and Optimized",
                 "Calibrate from MR. Must not hide intended A4/FreqPri.",
                 "UE may skip inter-frequency meas while serving RSRP is above SMeasure", "Silent A4 suppression", "Mobility Management in Connected Mode §4.1.5"]),
        ("row", ["6", "CELLHOPARACFG", "EutranFilterCoeffRsrp", "Basic and Optimized",
                 "Calibrate from MR. Over-smooth delays coverage rescue.",
                 "L3 RSRP filter", "Adds delay on top of TTT", "Mobility Management in Connected Mode Table 4-14"]),
        ("row", ["7", "INTERFREQHOGROUP", "InterFreqHoA1A2Hyst", "Basic and Optimized",
                 "Calibrate from MR. Keep A1/A2 hyst consistent.",
                 "A1/A2 hysteresis", "RSRP recommended (Table 4-15)", "Mobility Management in Connected Mode Table 4-9"]),
        ("row", ["8", "INTERFREQHOGROUP", "InterFreqHoA1A2TimeToTrig", "Basic and Optimized",
                 "Calibrate from MR",
                 "A1/A2 time-to-trigger", "QCI-specific optional", "Mobility Management in Connected Mode Table 4-9"]),
        ("row", ["9", "INTERFREQHOGROUP", "Coverage A2 threshold (correct family)", "Core parameter",
                 "Calibrate from MR. Book examples −85/−87 dBm are not design values.",
                 "Starts inter-frequency coverage measurement", "Separate families for A3 vs A4/A5 vs IRAT vs blind. Wrong family = wrong HO.", "Mobility Management in Connected Mode Tables 5-3, 5-10 ; §11.4.1.2"]),
        ("row", ["10", "INTERFREQHOGROUP", "A3 offset / hyst / TTT", "Basic and Optimized",
                 "Calibrate from MR",
                 "Relative HO among similar coverage", "Ofn/Ocn apply", "Mobility Management in Connected Mode Table 5-16"]),
        ("row", ["11", "INTERFREQHOGROUP", "InterFreqLoadBasedHoA4ThdRsrp", "Basic and Optimized",
                 "Calibrate from MR. Must be better than coverage A2. Book example −103 dBm is not a design value.",
                 "MLB/FreqPri absolute target gate", "Target need only be good enough, not better than serving", "Mobility Management in Connected Mode Table 5-22 ; Table 11-5 ; §11.4.1.2"]),
        ("row", ["12", "INTERFREQHOGROUP", "InterFreqHoA4Hyst", "Basic and Optimized",
                 "Calibrate from MR",
                 "A4 hysteresis", "Keep consistent with A4 thd", "Mobility Management in Connected Mode Table 11-5"]),
        ("row", ["13", "INTERFREQHOGROUP", "InterFreqHoA4TimeToTrig", "Core parameter",
                 "Must not be 5120 ms if FreqPri or MLB A4 is required",
                 "A4 time-to-trigger", "5120 ms disables FreqPri, CQI and service-based inter-frequency HO", "Mobility Management in Connected Mode Table 4-9 p.48"]),
        ("row", ["14", "INTERFREQHOGROUP", "A5 Thd1 / Thd2", "Basic and Optimized",
                 "Calibrate from MR",
                 "Serving poor AND target good", "Strongest coverage semantics. MLB A5 needs extra license.", "Mobility Management in Connected Mode Tables 5-18/5-19"]),
        ("row", ["15", "EUTRANINTERNFREQ", "MlbInterFreqHoEventType", "Core parameter",
                 "A4 for co-sited FDD. A5 only with Intra-LTE Load Balancing for Non-cosited Cells license.",
                 "Which event MLB uses to execute load HO", "Connected book is the HO engine; MLB book is the load brain", "Intra-RAT MLB Table 6-3 ; §6.1.1.5.2"]),
        ("row", ["16", "EUTRANINTERNFREQ", "IfMlbThdRsrpOffset", "Basic and Optimized",
                 "Per-frequency A4 offset. Calibrate from MR.",
                 "Shifts MLB A4 per EARFCN", "Plus operator/QCI offset if used", "Mobility Management in Connected Mode Table 11-5 ; Intra-RAT MLB p.137"]),
        ("row", ["17", "EUTRANINTERNFREQ", "FreqPriHoA4ThldRsrpOffset", "Basic and Optimized",
                 "Per-frequency FreqPri A4 offset. Calibrate from MR.",
                 "Shifts FreqPri A4 per EARFCN", "No reverse MLB target on a FreqPri pair", "Mobility Management in Connected Mode Table 11-5 ; p.303"]),
        ("row", ["18", "EUTRANINTERFREQNCELL", "CellIndividualOffset", "Basic and Optimized",
                 "Connected CIO (Ocn). Large CIO can mask RF overshoot.",
                 "Per-neighbour connected offset", "Fix RF first", "Mobility Management in Connected Mode pp.46–48"]),
        ("row", ["19", "FreqPri related MO", "MlbBasedFreqPriHoSwitch", "Core parameter",
                 "ON when MLB should own heavy-load decisions",
                 "Lets MLB override FreqPri under load", "Confirm exact MO name in MAE", "Mobility Management in Connected Mode Table 11-7"]),
        ("row", ["20", "FreqPri related MO", "LoadTriggerFreqPriHoSwitch", "Basic and Optimized",
                 "ON only if overlap, load info, neighbour not UE-number MLB triggered, no PCI conflict",
                 "Load-triggered FreqPri", "Waiting timer can wait for highest-priority freq", "Mobility Management in Connected Mode pp.302–303"]),
        ("row", ["21", "INTRARATHOCOMM", "FreqPriInHoProtectionTimer", "Basic and Optimized",
                 "Non-zero after incoming unnecessary HO",
                 "Stops immediate FreqPri bounce-back", "Ping-pong guard", "Mobility Management in Connected Mode p.300"]),
        ("row", ["22", "ENODEBALGOSWITCH", "AutoGapSwitch / GapPatternType", "Basic and Optimized",
                 "Acceptable TTI cost",
                 "Measurement gap", "Gaps steal DL TTIs (old UE / VoLTE)", "Mobility Management in Connected Mode Fig 4-10"]),
        ("space", 8),
        ("section", "How Connected locks to the other two books"),
        ("heads", H_STEP),
        ("row", ["← Idle", "Dedicated prio already discarded", "Idle only set the serving cell of this session", "Align idle thresholds with A1/A5", "Sheet 2", "Idle §5.1.3.1"]),
        ("row", ["→ MLB", "This book is the HO engine; MLB is the load brain", "ONLY_STRONGEST_CELL and A4 live in the MLB book", "LOAD_COVERAGE_MEAS_DECOUPLE_SW if coverage meas already on", "Sheet 4", "MLB Table 6-5 ; p.139"]),
        ("space", 8),
        ("section", "Connected activation sequence"),
        ("heads", H_MML),
        ("row", ["0", "LST", "Dump INTERFREQHOGROUP, EUTRANINTERNFREQ, NCELL, CELLHOPARACFG, HOMEASCOMM", "Read-only", "Do not paste example −85/−87/−103 dBm", "§11.4.1.2"]),
        ("row", ["1", "EUTRANINTERNFREQ", "Meas flag on; forbid-meas off for required HO targets", "NRT exists", "Most silent A4 failures are flags/NRT, not dBm", "§4.1.4.1.2"]),
        ("row", ["2", "CELLUEMEASCONTROLCFG", "Object cap ≥ number of inter-frequency objects this cell must measure", "Frequency plan known", "Otherwise equal-prio freqs drop at random", "p.125"]),
        ("row", ["3", "INTERFREQHOGROUP", "Coverage A1/A2/A5 from MR, correct A2 family", "Coverage design first", "Rescue before offload", "Tables 5-3, 5-10"]),
        ("row", ["4", "INTERFREQHOGROUP", "A4 from MR; better than coverage A2; TTT not 5120 ms if FreqPri/MLB A4 is required", "Step 3 done", "Main absolute target gate", "Table 4-9"]),
        ("row", ["5", "EUTRANINTERNFREQ", "MlbInterFreqHoEventType = A4 (A5 only with non-cosited MLB license)", "MLB will be used", "Co-sited FDD typically A4", "MLB Table 6-3"]),
        ("row", ["6", "FreqPri coordination", "MlbBasedFreqPriHoSwitch / LoadTrigger / non-zero incoming protect timer; no reverse MLB pair", "Confirm exact MO name in MAE", "Stops FreqPri fighting MLB", "p.303"]),
        ("space", 8),
        ("mml", MML_CONN),
    ]


def mlb():
    return [
        ("section", "Document charts"),
        ("fig", "mlb_31", "Intra-RAT MLB  Fig 3-1  — load balancing procedure"),
        ("fig", "mlb_41", "Intra-RAT MLB  Fig 4-1  — equalisation vs offload"),
        ("fig", "mlb_idle_conn", "Intra-RAT MLB  Figs 4-4 / 4-5  — idle transfer vs connected transfer"),
        ("section", "Procedure chart"),
        ("flow", ["Eval load N/C", "Trigger thd + offset", "Admit target", "Select UEs", "A4/A5 HO or idle release", "Penalty / stop"],
         "MLB Fig 3-1. Equalisation vs offload: Fig 4-1. Idle vs connected transfer: Figs 4-4 / 4-5. Load = N/C. Difference = (Load_s − Load_t) / Load_s."),
        ("space", 6),
        ("section", "Step by step"),
        ("heads", H_STEP),
        ("row", ["1", "Choose equalisation or offload", "Peer load available or not", "Equalisation when source and target load are known. Offload if not.", "Equalisation is safer when X2 / intra-eNB load exists.", "Fig 4-1 §4.1"]),
        ("row", ["2", "Choose the load indicator", "UE number vs PRB vs HW/transport", "UE number ≈ short non-GBR. PRB ≈ GBR / sustained. HW/transport = Low/Med/High/OverLoad.", "PRB MLB does not move CA UEs and is burst-sensitive.", "Fig 4-2, Fig 4-3"]),
        ("row", ["3", "Compute user-number load", "Load = N/C", "ActiveUeBasedLoadEvalSw uses DL-buffer UEs. SpectralEffBasedLoadEvalSw uses PRB, scale, GBR, SE (refresh 1 min if ≥10 UL-sync UEs).", "Raw UE-count on unequal bandwidth can reduce DL throughput.", "§5.1.1 ; Table 5-5 ; §6.1.2.2"]),
        ("row", ["4", "Trigger connected UE-number MLB", "InterFreqMlbUeNumThd + MlbUeNumOffset + MlbTrigJudgePeriod", "Enter when N ≥ thd+offset for the whole judge period. Leave when N < thd.", "Mode UE_NUMBER_ONLY + SynchronizedUE.", "Table 6-2 ; p.128"]),
        ("row", ["5", "Admit the target cell / frequency", "MlbTargetInd, OverlapInd, NoHoFlag, HO SR, HW/transport", "WITHOUT_IDLE_MLB / WITHOUT_CONNECT_MLB block one mode. Low/Med HW+transport. Not punished.", "Failed pair is not a target — fix RF/HO first (NCellHoSuccRateThld).", "pp.27–29"]),
        ("row", ["6", "Select UEs", "CELLMLBUESEL + protect timers", "UL-sync, not emergency, SPID/QCI, optional ARP/PRB/MCS/SNR. ONLY_STRONGEST_CELL recommended.", "Otherwise the UE is coverage-HO’d back.", "Table 6-5 p.159"]),
        ("row", ["7", "Limit volume and pick frequency", "MlbMaxUeNum ; FAIR / PRIORITY / LOADPRIORITY", "min(needed delta, hyst, max UE). LoadTransferEnhSw changes multi-target math.", "5 s eval + MlbMaxUeNum ≥ 40 over-transfers.", "pp.134–137"]),
        ("row", ["8", "Execute connected transfer", "A4 or A5 HO (FDD); optional blind", "Redirection generally not used in FDD. A5 needs extra license.", "Blind only with known containment.", "§6.1.1.5.2"]),
        ("row", ["9", "Execute idle transfer (if enabled)", "RRC release + dedicated prio + T320", "Class order NG-RAN > E-UTRAN low-load > E-UTRAN high-load > UTRAN > GERAN.", "Affects the next session only.", "§5.1.1.5"]),
        ("row", ["10", "Stop / punish", "CellPunishPrdNum × eval period", "Target reject (no radio resource) punishes the pair.", "Do not lower HO-success threshold to force MLB.", "p.29"]),
        ("space", 8),
        ("section", "Huawei cautions"),
        ("heads", H_NOTE),
        ("row", ["1", "ActiveUe when bandwidths differ; SpectralEff when SE differs a lot (e.g. >30%)", "Otherwise N/C is unfair across 15 vs 20 MHz.", "Turn load-model switches on before chasing the UE-number threshold.", "Step 3", "Table 5-5"]),
        ("row", ["2", "ONLY_STRONGEST_CELL", "Recommended cell-pick strategy.", "Without it, coverage HO bounces the UE back.", "Step 6", "Table 6-5"]),
        ("row", ["3", "Do not use InterFreqLoadEvalPrd = 5 s with MlbMaxUeNum ≥ 40", "Over-transfer.", "Lower volume or lengthen eval.", "Step 7", "p.136"]),
        ("row", ["4", "PRB-usage MLB skips CA UEs; gain dies when CA penetration ≳ 60%", "No experience fairness; bursty.", "Use as a supplement, not the main user-experience equaliser.", "Step 2", "Fig 6-2 ; §6.5"]),
        ("row", ["5", "Smart neighbour thds: 7 days collect, then refresh every 7 days", "Aggressive first-week seeds: extra CPU/HO and up to 5% TP swing.", "Need 15-min counter subscription. Do not treat learned thds as daily manual MOD targets.", "Step 4", "pp.29–31, 162"]),
        ("row", ["6", "Fixed-proportion idle + user-number connected MLB = ping-pong", "Adaptive-proportion idle not recommended without Huawei support.", "Keep that pair off together.", "Step 9", "§5.4.2.2 ; §5.5.2.1"]),
        ("row", ["7", "MLB and energy saving change each other’s targets", "A sleeping cell must not remain an MLB / high idle target.", "Coordinate ES.", "Step 5", "pp.153–156"]),
        ("space", 8),
        ("section", "What MLB gives / where it stops"),
        ("heads", ["Type", "Item", "Document statement", "Condition", "User impact", "Source"]),
        ("row", ["Gives", "True load equalisation", "Uses source and target load to reduce the difference", "Load exchange available", "Safer than blind offload", "Fig 4-1"]),
        ("row", ["Gives", "Idle + connected pair", "Idle steers next access; connected moves now", "T320 vs HO", "Idle has less gap/HO cost", "Figs 4-4 / 4-5"]),
        ("row", ["Stops", "Cannot fix RF, coverage, or band-capability mismatch", "Gain falls when coverage, UE band or PLMN differ", "p.39, 141, 214", "Not a substitute for RF", "§6.1.2.2"]),
        ("row", ["Stops", "Blind offload has no target-load check", "Can overload a cell that only looks low-load", "pp.167, 231", "Leave blind off unless designed", "Offload path"]),
        ("space", 8),
        ("section", "Parameters in the same order as the procedure"),
        ("note", "One parameter per row. Core parameter = required to activate the function. Basic and Optimized = used to tune performance after activation. Parameter Value is from the Huawei feature book; if the book gives no default, calibrate from MR. MLB Ch.8 points to the parameter reference for full ranges (p.299)."),
        ("heads", H_PAR),
        ("row", ["1", "CELLALGOSWITCH", "InterFreqMlbSwitch", "Core parameter",
                 "ON (InterFreqMlbSwitch-1)",
                 "Master intra-RAT connected MLB bit", "License + NRT + MlbTargetInd first", "Intra-RAT MLB Table 6-2"]),
        ("row", ["2", "CELLALGOSWITCH", "InterFreqIdleMlbSwitch", "Core parameter",
                 "ON if idle transfer is required",
                 "Idle dedicated-priority transfer", "T320 + SIB5 NORMAL + idle-allowed MlbTargetInd", "Intra-RAT MLB §5.1.1.5"]),
        ("row", ["3", "CELLALGOSWITCH", "InterFreqBlindMlbSwitch", "Core parameter",
                 "OFF unless containment is proven",
                 "Blind MLB", "Higher access risk without load exchange", "Intra-RAT MLB p.136"]),
        ("row", ["4", "CELLMLB", "MlbTriggerMode", "Core parameter",
                 "UE_NUMBER_ONLY typical for user-experience equalisation",
                 "Which load indicator starts MLB", "PRB_ONLY skips CA UEs", "Intra-RAT MLB Table 6-2 ; §6.5"]),
        ("row", ["5", "CELLMLB", "InterFreqUeTrsfType", "Core parameter",
                 "SynchronizedUE (connected). IdleUE for idle. PrbMlbSynchronizedUE for PRB.",
                 "Which UEs are transferable", "Must match trigger mode", "Intra-RAT MLB Table 6-2"]),
        ("row", ["6", "CELLMLB", "ActiveUeBasedLoadEvalSw", "Core parameter",
                 "ON when MLB frequencies have different bandwidths (Huawei)",
                 "N in Load=N/C uses DL-buffer UEs", "Turn on before chasing UE-number thd", "Intra-RAT MLB Table 5-5"]),
        ("row", ["7", "CELLMLB", "SpectralEffBasedLoadEvalSw", "Core parameter",
                 "ON when SE differs significantly (e.g. >30%) (Huawei)",
                 "C includes measured SE (refresh 1 min if ≥10 UL-sync UEs)", "Raw UE-count on unequal BW can reduce DL TP", "Intra-RAT MLB Table 5-5 ; §6.1.2.2"]),
        ("row", ["8", "CELLMLB", "LoadTransferEnhSw", "Basic and Optimized",
                 "ON for multi-target / PRB-diff calculation",
                 "Changes multi-target transfer math", "Several candidate frequencies", "Intra-RAT MLB p.134"]),
        ("row", ["9", "CELLMLB", "CaUserLoadTransferSw", "Basic and Optimized",
                 "ON only if CA UEs must be transferable",
                 "Allows CA/PCC UEs in user-number MLB", "If OFF, CA UEs are filtered. Target CA capability condition applies.", "Intra-RAT MLB pp.129–136, 157"]),
        ("row", ["10", "CELLMLB", "MlbHoCellSelectStrategy", "Core parameter",
                 "ONLY_STRONGEST_CELL (Huawei recommended)",
                 "Which neighbour is chosen for load HO", "Otherwise coverage HO bounces the UE back", "Intra-RAT MLB Table 6-5 p.159"]),
        ("row", ["11", "CELLMLB", "FreqSelectStrategy", "Basic and Optimized",
                 "FAIRSTRATEGY / PRIORITYBASED / LOADPRIORITY",
                 "How the target frequency is picked", "PRIORITYBASED uses MlbFreqPriority + freq penalty", "Intra-RAT MLB pp.137, 213"]),
        ("row", ["12", "CELLMLB", "InterFreqMlbUeNumThd", "Basic and Optimized",
                 "Calibrate from live load after ActiveUe/SE are ON. Feature book does not give a network-wide default.",
                 "UE-number enter/leave threshold (leave when N < thd)", "Enter uses thd + MlbUeNumOffset", "Intra-RAT MLB p.128"]),
        ("row", ["13", "CELLMLB", "MlbUeNumOffset", "Basic and Optimized",
                 "Calibrate from live load",
                 "Hysteresis on the UE-number trigger", "Enter = thd+offset", "Intra-RAT MLB p.128"]),
        ("row", ["14", "CELLMLB", "InterFreqIdleMlbUeNumThd", "Basic and Optimized",
                 "Calibrate from live idle-user load",
                 "Idle MLB trigger", "IdleUE transfer type", "Intra-RAT MLB Table 5-2"]),
        ("row", ["15", "CELLMLB", "MlbMaxUeNum", "Basic and Optimized",
                 "Never combine ≥40 with InterFreqLoadEvalPrd=5 s",
                 "Maximum UEs transferred per eval", "Over-transfer warning p.136", "Intra-RAT MLB p.136"]),
        ("row", ["16", "CELLMLB", "MlbTrigJudgePeriod", "Basic and Optimized",
                 "Trigger must hold for the whole period",
                 "How long the overload must persist", "Stops false trigger on short spikes", "Intra-RAT MLB p.128"]),
        ("row", ["17", "CELLMLB", "InterFreqLoadEvalPrd", "Basic and Optimized",
                 "Do not use 5 s together with MlbMaxUeNum≥40",
                 "Load evaluation period", "Over-transfer if too short and volume too high", "Intra-RAT MLB p.136"]),
        ("row", ["18", "EUTRANINTERNFREQ", "MlbTargetInd", "Core parameter",
                 "ALLOWED (or WITHOUT_IDLE_MLB / WITHOUT_CONNECT_MLB)",
                 "Whether the frequency may be an MLB target", "Coverage HO is separate (NoHoFlag / OverlapInd)", "Intra-RAT MLB pp.28, 129"]),
        ("row", ["19", "EUTRANINTERNFREQ", "MlbInterFreqHoEventType", "Core parameter",
                 "A4 (A5 only with non-cosited MLB license)",
                 "Event used to execute connected load HO", "FDD normally measurement-based HO, not redirection", "Intra-RAT MLB Table 6-3 ; §6.1.1.5.2"]),
        ("row", ["20", "EUTRANINTERNFREQ", "IfMlbThdRsrpOffset", "Basic and Optimized",
                 "Calibrate from MR",
                 "Per-frequency A4 offset for MLB", "Connected A4 must still beat coverage A2", "Intra-RAT MLB p.137"]),
        ("row", ["21", "CELLMLB", "NCellHoSuccRateThld", "Basic and Optimized",
                 "Pair HO success must stay ≥ this thd",
                 "Admit target only if HO SR is healthy", "Do not lower this to force MLB — fix RF/HO first", "Intra-RAT MLB p.27"]),
        ("row", ["22", "CELLMLB", "CellPunishPrdNum", "Basic and Optimized",
                 "Punish duration = this × eval period after target reject",
                 "Penalty after no-radio-resource reject", "Do not treat punish as an RF-threshold problem", "Intra-RAT MLB p.29"]),
        ("row", ["23", "CELLMLB", "MlbHoInProtectTimer", "Basic and Optimized",
                 "Non-zero re-MLB protect after incoming load HO",
                 "Ping-pong guard", "Pair with UE-select punish timer", "Intra-RAT MLB pp.130–135"]),
        ("row", ["24", "RRCCONNSTATETIMER", "T320ForLoadBalance", "Basic and Optimized",
                 "Idle dedicated-priority lifetime. SPID/PCC always 180 min.",
                 "Idle MLB release path", "Next session only", "Intra-RAT MLB §5.1.1.5"]),
        ("row", ["25", "EnhancedMlbAlgoSwitch", "DediPrioManageOnLowLoadSw", "Basic and Optimized",
                 "ON to hold released UEs off higher-load frequencies (verify license)",
                 "Idle dedicated-priority hold", "Table 5-10", "Intra-RAT MLB §5.3.1"]),
        ("row", ["26", "CELLMLB", "NCellTrigThldSmartOptAlgoSw", "Basic and Optimized",
                 "7 days collect, then calculate, refresh every 7 days. Needs 15-min counter subscription.",
                 "Learned pair thresholds", "First-week seeds: extra CPU/HO and up to 5% TP swing. Not a daily manual MOD target.", "Intra-RAT MLB pp.29–31, 162"]),
        ("space", 8),
        ("section", "How MLB locks to the other two books"),
        ("heads", H_STEP),
        ("row", ["← Idle", "Idle method = dedicated prio + T320 + SIB5 NORMAL", "MlbTargetInd can forbid idle targeting without blocking coverage HO", "UNDELIVER cannot be an idle target", "Sheet 2", "§5.1.1.5"]),
        ("row", ["← Connected", "HO engine A4/A5, meas flags, object cap, SMeasure, admission", "Trigger with 0 meas success = Connected sheet, not a thd problem", "ONLY_STRONGEST_CELL + A4 typical co-sited", "Sheet 3", "§6.1.1.5.2"]),
        ("space", 8),
        ("section", "MLB activation sequence"),
        ("heads", H_MML),
        ("row", ["0", "LST + license", "Dump CELLALGOSWITCH, CELLMLB, EUTRANINTERNFREQ, CELLMLBUESEL. Confirm Intra-RAT MLB license.", "MAE + parameter reference (Ch.8 p.299)", "Book does not list every default", "p.299"]),
        ("row", ["1", "EUTRANINTERNFREQ", "MlbTargetInd=ALLOWED, event A4, overlap and PERMIT_HO already valid", "Connected meas flags already correct", "WITHOUT_CONNECT_MLB / WITHOUT_IDLE_MLB to block one mode", "pp.28, 129"]),
        ("row", ["2", "Load model", "ActiveUe / SpectralEff / LoadTransferEnh as required (unequal BW / SE e.g. >30% / multi-target)", "Confirm bit names in MAE", "Do this before chasing UE-number thd", "Table 5-5"]),
        ("row", ["3", "CELLMLB", "UE_NUMBER_ONLY, SynchronizedUE, ONLY_STRONGEST_CELL, frequency strategy", "Targets already allowed", "Table 6-5 recommended pick", "Table 6-2"]),
        ("row", ["4", "CELLMLB", "UE-number thd/offset, MaxUeNum, judge/eval periods from live load — not 5 s + MaxUe≥40", "Load model ON", "Enter = thd+offset; leave = thd", "p.128, p.136"]),
        ("row", ["5", "CA / UE pick", "CaUserLoadTransferSw only if CA UEs must move. Then CELLMLBUESEL policy.", "CA plan", "If OFF, CA UEs stay filtered", "p.157"]),
        ("row", ["6", "Master bits last", "InterFreqMlbSwitch-1 and IdleMlbSwitch-1 if idle transfer is required. Blind stays 0.", "Steps 1–4 done. Idle also needs T320 and SIB5 NORMAL.", "Activate after targets exist", "Table 6-2"]),
        ("row", ["7", "Verify", "Load HO counters (Load vs UeNumLoad), HighLoad Dur/Num, meas success, DL active users, throughput, idle DedicatedPri. SON: Inter-Frequency Handover Statistics ; Idle Mode Release Statistics.", "15-min counters if smart thd is used", "When both PRB and UE-number MLB are on: PRB HO ≈ Load − UeNumLoad", "Tables 6-6, 6-21 ; §§6.1.4.2, 6.5.4.2"]),
        ("space", 8),
        ("mml", MML_MLB),
    ]


def activate():
    return [
        ("section", "One sequence across all three books"),
        ("note", "Do Idle camping first (next access layer), then Connected coverage and A4 (HO engine), then MLB load brain last. LST before every MOD. Placeholders: LocalCellId=<x>, DlEarfcn=<earfcn>. Confirm enums in MAE. Example dBm in the Connected book are not design values."),
        ("fig", "chain", "Activation follows the same chain: Idle Fig 5-1 → Connected Fig 4-1 → MLB Fig 3-1"),
        ("flow", ["LST baseline", "Idle SIB / priority / search", "Connected flags + coverage", "A4 gate from MR", "MLB target + load model", "MLB thd + master bit + T320"],
         "Activation order — not the sample-file MML"),
        ("space", 6),
        ("heads", ["Step", "Book", "MO", "What to do", "Stop if this is missing", "Source"]),
        ("row", ["0", "All", "LST *", "Dump current CELLRESEL, EUTRANINTERNFREQ, NCELL, INTERFREQHOGROUP, CELLHOPARACFG, HOMEASCOMM, CELLALGOSWITCH, CELLMLB, CELLMLBUESEL, RRCCONNSTATETIMER. Check Intra-RAT MLB license.", "No baseline = no rollback", "MAE"]),
        ("row", ["1", "Idle", "CELLRESEL", "Serving common priority. CFG search. SIntraSearch > SNonIntraSearch (example SNonIntraSearch=10). Capacity/hotspot ranks above coverage layer.", "SIB3 not planned", "Idle §5.4.1.1 ; MLB §5.1.2.1"]),
        ("row", ["2", "Idle", "EUTRANINTERNFREQ", "Every needed EARFCN: CfgInd=CFG, priority, NORMAL (not UNDELIVER), reselection thresholds from MR, MlbTargetInd as designed.", "Frequency missing from SIB5 plan", "Idle §5.1.3.1"]),
        ("row", ["3", "Idle", "CELLRESEL", "ThrshServLow aligned with later coverage A2/A5 thinking.", "Coverage A2 not yet known — set after step 5 or use current A2", "Idle Tables 5-3 / 5-4"]),
        ("row", ["4", "Connected", "EUTRANINTERNFREQ + NCELL", "Meas flag on; forbid-meas off for HO targets; symmetric NRT; NoHoFlag PERMIT; no PCI conflict.", "Most silent A4 failures start here", "Connected §4.1.4"]),
        ("row", ["5", "Connected", "CELLUEMEASCONTROLCFG + HOMEASCOMM", "Object cap ≥ needed freqs. SMeasure must not hide A4.", "Equal-prio random drop / silent A4", "p.125 ; §4.1.5"]),
        ("row", ["6", "Connected", "INTERFREQHOGROUP", "Coverage A1/A2/A5 from MR, correct A2 family. Then A4 better than that A2. A4 TTT ≠ 5120 ms if FreqPri or MLB A4 is required.", "Do not paste −85/−87/−103 dBm from the book", "§11.4.1.2 ; Table 4-9"]),
        ("row", ["7", "Connected", "FreqPri coordination", "If FreqPri is used with MLB: MlbBasedFreqPriHoSwitch / LoadTrigger as designed; incoming protect timer non-zero; no reverse MLB pair.", "Ping-pong A↔B", "p.303"]),
        ("row", ["8", "MLB", "EUTRANINTERNFREQ", "MlbTargetInd and MlbInterFreqHoEventType=A4 (A5 only with non-cosited MLB license).", "Overlap / PERMIT_HO / meas already done in steps 4–6", "MLB pp.28, 129 ; Table 6-3"]),
        ("row", ["9", "MLB", "Load-model switches", "ActiveUe if bandwidths differ. SpectralEff if SE differs a lot (e.g. >30%). LoadTransferEnh if multi-target.", "Do this before the UE-number threshold", "Table 5-5"]),
        ("row", ["10", "MLB", "CELLMLB", "UE_NUMBER_ONLY, SynchronizedUE, ONLY_STRONGEST_CELL, frequency strategy, thd/offset/MaxUeNum/periods. Never 5 s eval + MaxUe ≥ 40.", "Load model from step 9", "Table 6-2, 6-5 ; p.136"]),
        ("row", ["11", "MLB + Idle", "CELLALGOSWITCH + T320", "Master MLB bits last. Idle bit only if SIB5 NORMAL and T320 set. Blind bit stays 0 unless containment is proven.", "License + steps 1–10", "Table 6-2"]),
        ("row", ["12", "All", "Verify after SI delay", "Idle: wait next SI modification period; dedicated-pri counters. Connected: Coverage / FreqPri HO SR. MLB: Load HO, meas success, HighLoad, DL TP, SON handover / idle-release logs.", "Same 15-min as the change is too early for idle", "Idle §7.1.3 ; MLB Tables 6-6, 6-21"]),
        ("space", 8),
        ("section", "License reminders (confirm exact names in MAE)"),
        ("heads", ["No.", "License / package", "What it enables", "If missing", "Where used", "Source"]),
        ("row", ["1", "Basic idle + coverage A1–A5", "Usually in the LTE eNodeB package", "If missing, nothing else in this file will look right", "Sheets 2–3", "—"]),
        ("row", ["2", "Intra-RAT Mobility Load Balancing", "Connected and idle MLB bits", "Switch ON but Load HO / DedicatedPri counters stay 0", "Sheet 4", "MLB Table 6-2"]),
        ("row", ["3", "Intra-LTE Load Balancing for Non-cosited Cells", "Required for MLB event A5", "Stay on A4 without it", "Sheet 3–4", "MLB pp.142–143"]),
        ("row", ["4", "FreqPri / service-based mobility package", "A1 high-band steering", "Verify the sold name", "Sheet 3", "Ch.11"]),
        ("row", ["5", "Blind / CA-transfer / SON smart-thd options", "Sold separately in some packages", "Leave off if the license is not present", "Sheet 4", "MLB"]),
        ("space", 8),
        ("major", "SymbolShutdownSwitch in the old CSV sample is Symbol Power Saving. It is not a mobility command and is not used here."),
        ("space", 8),
        ("mml", MML_ALL),
    ]


def main():
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    figs = build_figures()
    wb = Workbook()
    wb.active.title = "tmp"
    build_sheet(wb, "Read me", "4G LTE Mobility Management — eRAN21.1 v3.3  |  How to read this file", overview(), BLUE, figs)
    build_sheet(wb, "1. End-to-end chain", "1. End-to-end chain  |  Idle → Connected → MLB → Idle", chain(), "1F4E79", figs)
    build_sheet(wb, "2. Idle Mode", "2. Idle Mode Management  |  eRAN21.1 Issue 04  |  Step by step", idle(), "008000", figs)
    build_sheet(wb, "3. Connected Mode", "3. Mobility Management in Connected Mode  |  eRAN21.1 Issue 08  |  Step by step", connected(), "2E75B6", figs)
    build_sheet(wb, "4. Intra-RAT MLB", "4. Intra-RAT Mobility Load Balancing  |  eRAN21.1 Issue 10  |  Step by step", mlb(), "C65911", figs)
    build_sheet(wb, "5. Activation order", "5. Activation order  |  One sequence across the three books", activate(), BLUE, figs)
    del wb["tmp"]
    wb.properties.title = "4G LTE Mobility Management eRAN21.1 v3.3"
    wb.properties.subject = "Step-by-step summary with document charts and Combined MML"
    wb.properties.version = "3.3"
    wb.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
